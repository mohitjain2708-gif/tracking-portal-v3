from __future__ import annotations

import json
import logging
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta
from difflib import SequenceMatcher
from html import unescape
from io import BytesIO
from pathlib import Path
from threading import Lock, Thread
from typing import Any
from urllib.parse import quote
from uuid import uuid4
from zoneinfo import ZoneInfo

import boto3
import requests
from botocore.exceptions import BotoCoreError, ClientError
from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from jose import JWTError, jwt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import create_engine, func, inspect, select, text
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import database_url as normalized_database_url, engine, get_db
from app.core.security import hash_password
from app.models.customer_directory import CustomerDirectory
from app.models.audit_log import AuditLog
from app.models.shipment import Shipment
from app.models.shipment_source import RawSourceRow, ShipmentBatch, ShipmentSource, SourceMappingProfile, SourceConnection
from app.models.upload_session import UploadSession
from app.models.user import User
from app.schemas.shipment import ShipmentCreateRequest, ShipmentGroupUpdateRequest, ShipmentStatusUpdateRequest
from app.services.shipment_state import (
    MAX_SAME_CYCLE_BACKWARD_DRIFT_DAYS,
    STALE_SOURCE_DAYS,
    extract_concor_candidates,
    extract_ldb_candidates,
    extract_pristine_candidates,
    pristine_birgunj_can_override_lagging_live_feeds,
    resolve_shipment_state,
)

router = APIRouter()
logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parents[3]
RUNTIME_DIR = Path(settings.runtime_dir).expanduser()
if not RUNTIME_DIR.is_absolute():
    RUNTIME_DIR = (BASE_DIR / RUNTIME_DIR).resolve()
TEMP_IMPORT_DIR = RUNTIME_DIR / "temp_imports"
STATE_FILE = RUNTIME_DIR / "shipment_state.json"
CACHE_FILE = RUNTIME_DIR / "tracking_cache.json"
REFRESH_JOBS_FILE = RUNTIME_DIR / "refresh_jobs.json"
BACKGROUND_REFRESH_LEASE_FILE = RUNTIME_DIR / "background_refresh_lease.json"
BL_DOCUMENTS_DIR = RUNTIME_DIR / "bl_documents"
BL_DOCUMENTS_INDEX_FILE = RUNTIME_DIR / "bl_documents_index.json"
LOCATION_DISTANCE_CACHE_FILE = RUNTIME_DIR / "location_distance_cache.json"
DEFAULT_LOCAL_USER_EMAIL = settings.demo_email
APP_TIMEZONE = ZoneInfo("Asia/Kolkata")
REFRESH_JOB_RETENTION_HOURS = 12
REFRESH_JOB_STALE_SECONDS = 600
BACKGROUND_REFRESH_LEASE_SECONDS = 900
STALE_PORTAL_RECORD_DAYS = STALE_SOURCE_DAYS

RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
TEMP_IMPORT_DIR.mkdir(parents=True, exist_ok=True)
BL_DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)

LDB_API_URL = "https://www.ldb.co.in/api/ldb/container/search"
CONCOR_API_URL = "https://www.concorindia.co.in/api/multipalContainer"
PRISTINE_TRACKING_URL = "https://pristinevalleyport.com/tracking/"
CACHE_DURATION_MINUTES = 30
TRACKING_SOURCE_TIMEOUT_SECONDS = 10
TRACKING_POOL_WORKERS = 6
NOMINATIM_SEARCH_URL = "https://nominatim.openstreetmap.org/search"
OSRM_ROUTE_URL = "https://router.project-osrm.org/route/v1/driving"
BIRGUNJ_REFERENCE = {"lat": 27.0104, "lon": 84.8774}
LOCATION_COORDINATE_FALLBACKS = [
    {"patterns": ["BIRGUNJ", "BIRGANJ"], "lat": 27.0104, "lon": 84.8774},
    {"patterns": ["RAXAUL"], "lat": 26.9795, "lon": 84.8507},
    {"patterns": ["MUZAFFARPUR"], "lat": 26.1209, "lon": 85.3647},
    {"patterns": ["SONPUR"], "lat": 25.6992, "lon": 85.1949},
    {"patterns": ["PATNA"], "lat": 25.5941, "lon": 85.1376},
    {"patterns": ["KHURDA ROAD", "KHORDHA ROAD"], "lat": 20.1535, "lon": 85.7086},
    {"patterns": ["BANSPANI"], "lat": 21.6309, "lon": 85.5850},
    {"patterns": ["DANGOAPOSI"], "lat": 22.0644, "lon": 85.6432},
    {"patterns": ["TATANAGAR", "JAMSHEDPUR"], "lat": 22.8046, "lon": 86.2029},
    {"patterns": ["KOLKATA", "SYAMA PRASAD", "SMP"], "lat": 22.5726, "lon": 88.3639},
    {"patterns": ["HALDIA", "HICT"], "lat": 22.0257, "lon": 88.0583},
    {"patterns": ["VISHAKAPATNAM", "VISAKHAPATNAM", "VIZAG", "MMLP VISHAKAPATNAM"], "lat": 17.6868, "lon": 83.2185},
]
VALID_DOCUMENT_TYPES = ("invoice", "packing_list", "bl_copy")
SHIPMENT_STATUS_PRIORITY = {"active": 3, "completed": 2, "archived": 1}
MOVEMENT_PRIORITY = {"Arrived Birgunj": 4, "On Rail": 3, "At Port": 2, "Hi Seas": 1}
BUSINESS_SUFFIXES = {
    "PVT",
    "PVT.",
    "LTD",
    "LTD.",
    "LLP",
    "LLC",
    "INC",
    "INC.",
    "CO",
    "CO.",
    "PLC",
}
CUSTOMER_IGNORE_TOKENS = BUSINESS_SUFFIXES | {
    "PRIVATE",
    "LIMITED",
    "ENTERPRISE",
    "ENTERPRISES",
    "INTERNATIONAL",
    "TRADING",
}

SHIPMENT_COLUMN_DEFINITIONS = {
    "latest_location": "TEXT NOT NULL DEFAULT ''",
    "latest_time": "VARCHAR(32) NOT NULL DEFAULT ''",
    "port_arrival_date": "VARCHAR(32) NOT NULL DEFAULT ''",
    "birgunj_arrival_date": "VARCHAR(32) NOT NULL DEFAULT ''",
    "pristine_booking_date": "VARCHAR(32) NOT NULL DEFAULT ''",
    "train_no": "VARCHAR(64) NOT NULL DEFAULT ''",
    "departure": "VARCHAR(32) NOT NULL DEFAULT ''",
    "wagon_loaded_date": "VARCHAR(32) NOT NULL DEFAULT ''",
    "concor_location_code": "VARCHAR(32) NOT NULL DEFAULT ''",
    "rail_status": "VARCHAR(64) NOT NULL DEFAULT ''",
    "movement_category": "VARCHAR(64) NOT NULL DEFAULT 'Hi Seas'",
    "delay_days": "FLOAT NOT NULL DEFAULT 0",
    "wagon_no": "VARCHAR(64) NOT NULL DEFAULT ''",
    "train_origin": "TEXT NOT NULL DEFAULT ''",
    "train_destination": "TEXT NOT NULL DEFAULT ''",
    "shipping_line": "VARCHAR(255) NOT NULL DEFAULT ''",
    "tracking_source": "VARCHAR(64) NOT NULL DEFAULT ''",
    "last_refresh_at": "VARCHAR(32) NOT NULL DEFAULT ''",
    "last_refresh_status": "VARCHAR(64) NOT NULL DEFAULT ''",
    "last_refresh_error": "TEXT NOT NULL DEFAULT ''",
    "clearance_doc_number": "VARCHAR(120) NOT NULL DEFAULT ''",
    "do_date": "VARCHAR(32) NOT NULL DEFAULT ''",
    "document_status": "VARCHAR(32) NOT NULL DEFAULT ''",
    "original_docs_received_date": "VARCHAR(32) NOT NULL DEFAULT ''",
    "source_type": "VARCHAR(64) NOT NULL DEFAULT 'manual'",
    "source_label": "VARCHAR(255) NOT NULL DEFAULT 'Manual Entry'",
    "source_batch_id": "INTEGER NOT NULL DEFAULT 0",
    "raw_source_row_id": "INTEGER NOT NULL DEFAULT 0",
}

PORT_PATTERNS = [
    "VISHAKAPATNAM", "VISAKHAPATNAM", "MMLP VISHAKAPATNAM", "MMLP/VISHAKAPATNAM",
    "MMLP-VISHAKAPATNAM", "KOLKATA/SYAMA PRASAD MOOKERJEE PORT", "SYAMA PRASAD MOOKERJEE PORT",
    "SMP, KOLKATA", "HALDIA INTERNATIONAL CONTAINER TERMINAL", "HICT",
    "PURBA MEDINIPUR/HALDIA INTERNATIONAL CONTAINER TERMINAL", "VISAKHA CONTAINER TERMINAL",
    "VPL INTEGRAL CFS", "KOLKATA", "HALDIA", "PORT", "CFS",
]

_storage_ready = False
_refresh_jobs: dict[str, dict[str, Any]] = {}
_refresh_jobs_lock = Lock()
_background_refresh_instance_id = uuid4().hex
_documents_index_cache: dict[str, dict[str, Any]] | None = None
_documents_index_mtime: float | None = None
_r2_client = None


def _now_datetime() -> str:
    return datetime.now(APP_TIMEZONE).strftime("%d-%m-%Y %H:%M:%S")


def _now_iso_for_cache() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _format_app_timestamp(value: datetime | None) -> str:
    if not value:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=ZoneInfo("UTC"))
    return value.astimezone(APP_TIMEZONE).strftime("%d-%m-%Y %H:%M:%S")


def _load_state() -> dict[str, Any]:
    if not STATE_FILE.exists():
        return {"shipments": [], "shipment_id_counter": 1}
    try:
        data = json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, TypeError) as exc:
        logger.warning("Unable to read shipment state from %s: %s", STATE_FILE, exc)
        return {"shipments": [], "shipment_id_counter": 1}
    data.setdefault("shipments", [])
    data.setdefault("shipment_id_counter", 1)
    return data


def _parse_refresh_job_timestamp(value: Any) -> datetime | None:
    text = _clean_text(value)
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=APP_TIMEZONE)
    return parsed


def _load_refresh_jobs() -> dict[str, dict[str, Any]]:
    if not REFRESH_JOBS_FILE.exists():
        return {}
    try:
        payload = json.loads(REFRESH_JOBS_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, TypeError) as exc:
        logger.warning("Unable to read refresh jobs from %s: %s", REFRESH_JOBS_FILE, exc)
        return {}
    return payload if isinstance(payload, dict) else {}


def _save_refresh_jobs(jobs: dict[str, dict[str, Any]]) -> None:
    REFRESH_JOBS_FILE.write_text(json.dumps(jobs, indent=2, ensure_ascii=True), encoding="utf-8")


def _load_background_refresh_lease() -> dict[str, Any]:
    if not BACKGROUND_REFRESH_LEASE_FILE.exists():
        return {}
    try:
        payload = json.loads(BACKGROUND_REFRESH_LEASE_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, TypeError) as exc:
        logger.warning(
            "Unable to read background refresh lease from %s: %s",
            BACKGROUND_REFRESH_LEASE_FILE,
            exc,
        )
        return {}
    return payload if isinstance(payload, dict) else {}


def _save_background_refresh_lease(payload: dict[str, Any]) -> None:
    BACKGROUND_REFRESH_LEASE_FILE.write_text(
        json.dumps(payload, indent=2, ensure_ascii=True),
        encoding="utf-8",
    )


def _update_background_refresh_lease(**extra: Any) -> None:
    payload = {
        "instance_id": _background_refresh_instance_id,
        "heartbeat_at": datetime.now(APP_TIMEZONE).isoformat(),
    }
    payload.update(extra)
    _save_background_refresh_lease(payload)


def _try_acquire_background_refresh_lease() -> bool:
    lease = _load_background_refresh_lease()
    heartbeat = _parse_refresh_job_timestamp(lease.get("heartbeat_at"))
    now = datetime.now(APP_TIMEZONE)
    if (
        lease
        and lease.get("instance_id") != _background_refresh_instance_id
        and heartbeat
        and (now - heartbeat).total_seconds() < BACKGROUND_REFRESH_LEASE_SECONDS
    ):
        return False
    _update_background_refresh_lease(state="running")
    return True


def _cleanup_refresh_jobs(jobs: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    cutoff = datetime.now(APP_TIMEZONE) - timedelta(hours=REFRESH_JOB_RETENTION_HOURS)
    cleaned: dict[str, dict[str, Any]] = {}
    for task_id, job in jobs.items():
        if not isinstance(job, dict):
            continue
        updated_at = _parse_refresh_job_timestamp(job.get("updated_at"))
        if updated_at and updated_at < cutoff:
            continue
        cleaned[task_id] = job
    return cleaned


def _normalize_refresh_job(job: dict[str, Any] | None) -> dict[str, Any] | None:
    if not isinstance(job, dict):
        return None
    normalized = dict(job)
    updated_at = _parse_refresh_job_timestamp(normalized.get("updated_at"))
    state = _clean_text(normalized.get("state")).lower()
    if updated_at and state in {"queued", "running"}:
        age_seconds = (datetime.now(APP_TIMEZONE) - updated_at).total_seconds()
        if age_seconds > REFRESH_JOB_STALE_SECONDS:
            normalized["state"] = "failed"
            normalized["progress"] = 100
            normalized["message"] = "Live tracking refresh took too long and was stopped. Please try again."
            normalized["updated_at"] = datetime.now(APP_TIMEZONE).isoformat()
    return normalized


def _clean_text(value: Any) -> str:
    return str(value or "").strip()


def _sheet_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).strip()
    return str(value).strip()


def _clean_container(value: Any) -> str:
    return str(value or "").strip().upper()


def _clean_bl(value: Any) -> str:
    return _sheet_cell_text(value).upper()


def _normalize_bl_number(value: Any) -> str:
    normalized = re.sub(r"\s+", "", _clean_bl(value))
    if "." in normalized:
        whole, fraction = normalized.rsplit(".", 1)
        if whole and fraction and set(fraction) == {"0"} and re.fullmatch(r"[A-Z0-9/-]+", whole):
            return whole
    return normalized


def _container_format_error(container_number: str) -> str | None:
    normalized = _clean_container(container_number)
    if not normalized:
        return "Container number is required."
    if not re.fullmatch(r"[A-Z]{4}\d{7}", normalized):
        return f"{normalized} must use 4 letters followed by 7 digits."
    return None


def _parse_container_numbers(values: list[Any] | None = None, fallback: Any = "") -> tuple[list[str], list[str]]:
    raw_values: list[Any] = list(values or [])
    if not raw_values and fallback not in (None, ""):
        raw_values = [fallback]

    tokens: list[str] = []
    for value in raw_values:
        text_value = _clean_text(value)
        if not text_value:
            continue
        parts = re.split(r"[\s,;]+", text_value)
        tokens.extend(part for part in parts if _clean_text(part))

    normalized: list[str] = []
    invalid: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        container_number = _clean_container(token)
        if not container_number or container_number in seen:
            continue
        seen.add(container_number)
        if _container_format_error(container_number):
            invalid.append(container_number)
        else:
            normalized.append(container_number)
    return normalized, invalid


def _containers_from_import_row(row_obj: dict[str, Any], container_col: str) -> tuple[list[str], list[str], str]:
    raw_value = row_obj.get(container_col, "")
    valid_containers, invalid_containers = _parse_container_numbers(fallback=raw_value)
    return valid_containers, invalid_containers, _clean_text(raw_value)


def _row_to_strings(row: list[Any]) -> list[str]:
    return [_clean_text(value) for value in row]


def _pick_header_row(rows: list[list[Any]]) -> tuple[int, list[str]]:
    best_idx = 0
    best_count = -1
    best_headers: list[str] = []
    for idx, row in enumerate(rows[:10]):
        cleaned = _row_to_strings(row)
        count = sum(1 for value in cleaned if value)
        if count > best_count:
            best_count = count
            best_idx = idx
            best_headers = cleaned
    headers = [header if header else f"Column {index + 1}" for index, header in enumerate(best_headers)]
    return best_idx + 1, headers


def _merged_cell_value_map(sheet) -> dict[tuple[int, int], Any]:
    merged_values: dict[tuple[int, int], Any] = {}
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left_value = sheet.cell(row=min_row, column=min_col).value
        for row_index in range(min_row, max_row + 1):
            for col_index in range(min_col, max_col + 1):
                merged_values[(row_index, col_index)] = top_left_value
    return merged_values


def _sheet_rows_with_merged_fill(sheet) -> list[list[Any]]:
    merged_values = _merged_cell_value_map(sheet)
    rows: list[list[Any]] = []
    for row in sheet.iter_rows():
        values: list[Any] = []
        for cell in row:
            values.append(merged_values.get((cell.row, cell.column), cell.value))
        rows.append(values)
    return rows


def _build_row_object(headers: list[str], values: list[Any]) -> dict[str, Any]:
    row_obj: dict[str, Any] = {}
    for index, header in enumerate(headers):
        row_obj[header] = _sheet_cell_text(values[index] if index < len(values) else "")
    return row_obj


def _normalize_import_row_objects(
    headers: list[str],
    rows: list[list[Any]],
    header_index: int,
    carry_columns: list[str],
) -> list[dict[str, Any]]:
    normalized_rows: list[dict[str, Any]] = []
    carry_forward = {column: "" for column in carry_columns if column}

    for offset, raw_row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        row_obj = _build_row_object(headers, list(raw_row))
        if not any(str(value).strip() for value in row_obj.values()):
            carry_forward = {column: "" for column in carry_forward}
            continue

        for column in carry_columns:
            if not column or column not in row_obj:
                continue
            current_value = _clean_text(row_obj.get(column))
            if current_value:
                carry_forward[column] = current_value
            elif carry_forward.get(column):
                row_obj[column] = carry_forward[column]

        row_obj["__source_row_number"] = offset
        normalized_rows.append(row_obj)

    return normalized_rows


def _apply_import_row_overrides(
    normalized_rows: list[dict[str, Any]],
    row_overrides: list[dict[str, Any]] | None,
    customer_col: str,
    container_col: str,
    bl_col: str,
) -> list[dict[str, Any]]:
    if not row_overrides:
        return normalized_rows

    overrides_by_row: dict[int, dict[str, Any]] = {}
    for override in row_overrides:
        try:
            row_number = int(override.get("source_row_number") or 0)
        except (TypeError, ValueError):
            row_number = 0
        if row_number <= 0:
            continue
        overrides_by_row[row_number] = override

    if not overrides_by_row:
        return normalized_rows

    updated_rows: list[dict[str, Any]] = []
    for row_obj in normalized_rows:
        row_number = int(row_obj.get("__source_row_number") or 0)
        override = overrides_by_row.get(row_number)
        if not override:
            updated_rows.append(row_obj)
            continue
        next_row = dict(row_obj)
        if customer_col:
            next_row[customer_col] = _clean_text(override.get("customer_name", next_row.get(customer_col, "")))
        if container_col:
            next_row[container_col] = _clean_text(override.get("container_number", next_row.get(container_col, "")))
        if bl_col:
            next_row[bl_col] = _clean_text(override.get("bl_number", next_row.get(bl_col, "")))
        updated_rows.append(next_row)
    return updated_rows


def _summarize_import_rows(
    normalized_rows: list[dict[str, Any]],
    customer_col: str,
    container_col: str,
    bl_col: str,
) -> dict[str, Any]:
    invalid_rows: list[dict[str, Any]] = []
    skipped_blank_count = 0
    valid_count = 0

    for row_obj in normalized_rows:
        source_row_number = int(row_obj.get("__source_row_number") or 0)
        customer_name = _clean_text(row_obj.get(customer_col))
        bl_number = _clean_text(row_obj.get(bl_col))
        valid_containers, invalid_containers, raw_container_value = _containers_from_import_row(row_obj, container_col)

        if not valid_containers and not invalid_containers:
            skipped_blank_count += 1
            continue

        if invalid_containers:
            invalid_rows.append(
                {
                    "source_row_number": source_row_number,
                    "customer_name": customer_name,
                    "container_number": raw_container_value,
                    "bl_number": bl_number,
                    "error": f"{' '.join(invalid_containers)} must use 4 letters followed by 7 digits.",
                }
            )
            continue

        valid_count += len(valid_containers)

    return {
        "valid_count": valid_count,
        "invalid_rows": invalid_rows,
        "invalid_count": len(invalid_rows),
        "skipped_blank_count": skipped_blank_count,
        "total_reviewable_rows": valid_count + len(invalid_rows),
    }


def _detect_import_duplicates(
    db: Session,
    current_user: User,
    normalized_rows: list[dict[str, Any]],
    customer_col: str,
    container_col: str,
    bl_col: str,
) -> dict[str, Any]:
    existing_keys = {
        (
            shipment.container_number,
            _normalize_bl_number(shipment.bl_number),
            _format_customer_name(shipment.customer_name),
        )
        for shipment in db.execute(
            _user_shipment_select(current_user).where(Shipment.shipment_status != "archived")
        ).scalars()
    }
    seen_in_import: set[tuple[str, str, str]] = set()
    duplicate_rows: list[dict[str, Any]] = []
    for row_obj in normalized_rows:
        source_row_number = int(row_obj.get("__source_row_number") or 0)
        customer_name = _format_customer_name(_clean_text(row_obj.get(customer_col)))
        bl_number = _normalize_bl_number(row_obj.get(bl_col))
        valid_containers, invalid_containers, raw_container_value = _containers_from_import_row(row_obj, container_col)
        if invalid_containers or (not valid_containers and not raw_container_value):
            continue
        for container_number in valid_containers:
            shipment_key = (container_number, bl_number, customer_name)
            if shipment_key in seen_in_import:
                duplicate_rows.append(
                    {
                        "source_row_number": source_row_number,
                        "customer_name": customer_name,
                        "container_number": container_number,
                        "bl_number": bl_number,
                        "error": "Duplicate shipment row repeated in this source",
                    }
                )
                continue
            seen_in_import.add(shipment_key)
            if shipment_key in existing_keys:
                duplicate_rows.append(
                    {
                        "source_row_number": source_row_number,
                        "customer_name": customer_name,
                        "container_number": container_number,
                        "bl_number": bl_number,
                        "error": "This shipment already exists in the portal",
                    }
                )
    return {
        "duplicate_rows": duplicate_rows,
        "duplicate_count": len(duplicate_rows),
    }


def _read_sheet(file_path: Path, preferred_sheet: str | None = None) -> tuple[str, int, list[str], list[dict[str, Any]], list[list[Any]], list[str]]:
    workbook = load_workbook(file_path, data_only=True)
    sheet = _select_sheet_by_name(workbook, preferred_sheet)
    rows = _sheet_rows_with_merged_fill(sheet)
    if not rows:
        raise HTTPException(status_code=400, detail="Excel file is empty")
    header_row, headers = _pick_header_row(rows)
    preview_rows: list[dict[str, Any]] = []
    for row in rows[header_row : header_row + 5]:
        record = _build_row_object(headers, list(row))
        if any(str(value).strip() for value in record.values()):
            preview_rows.append(record)
    return sheet.title, header_row, headers, preview_rows, rows, list(workbook.sheetnames)


def _score_sheet_relevance(sheet_name: str, rows: list[list[Any]]) -> tuple[int, int, int, int]:
    normalized_sheet_name = _clean_text(sheet_name).upper()
    header_row, headers = _pick_header_row(rows[:25] or rows)
    normalized_headers = [_clean_text(header).upper() for header in headers]
    container_header_score = sum(
        1
        for header in normalized_headers
        if any(keyword in header for keyword in ("CONTAINER", "CNTR", "BL", "B/L", "BILL OF LADING"))
    )
    data_row_score = sum(1 for row in rows[header_row:] if any(_clean_text(value) for value in row))
    non_empty_score = sum(1 for row in rows if any(_clean_text(value) for value in row))
    title_score = 0
    if data_row_score > 0:
        if "OONC" in normalized_sheet_name:
            title_score += 200
        if "TRACK" in normalized_sheet_name:
            title_score += 40
    return title_score, container_header_score, data_row_score, non_empty_score


def _select_relevant_sheet(workbook) -> Any:
    best_sheet = None
    best_score: tuple[int, int, int, int, int] | None = None
    for index, sheet_name in enumerate(workbook.sheetnames):
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        score = (*_score_sheet_relevance(sheet_name, rows), -index)
        if best_score is None or score > best_score:
            best_sheet = sheet
            best_score = score
    if best_sheet is None:
        raise HTTPException(status_code=400, detail="Excel file is empty")
    return best_sheet


def _select_sheet_by_name(workbook, sheet_name: str | None = None):
    requested_sheet = _clean_text(sheet_name)
    if requested_sheet:
        for candidate in workbook.sheetnames:
            if _clean_text(candidate).casefold() == requested_sheet.casefold():
                return workbook[candidate]
        raise HTTPException(status_code=400, detail=f"Worksheet '{requested_sheet}' was not found in the workbook")
    return _select_relevant_sheet(workbook)


def _is_arrived(location: str) -> bool:
    location_text = _clean_text(location).upper()
    return "BIRGUNJ" in location_text or "BIRGANJ" in location_text


def _is_port(location: str) -> bool:
    location_text = _clean_text(location).upper()
    return bool(location_text) and any(pattern in location_text for pattern in PORT_PATTERNS)


def _normalize_existing_movement(value: str) -> str:
    movement = _clean_text(value)
    if not movement or movement == "High Seas":
        return "Hi Seas"
    if movement in {"Arrived", "Arrived Birgunj"}:
        return "Arrived Birgunj"
    if movement in {"Moving", "In Transit", "On Rail"}:
        return "On Rail"
    if movement in {"At Origin", "Delayed", "At Port"}:
        return "At Port"
    return movement


def _movement_category(
    location: str,
    train_no: str,
    departure: str,
    delay_days: float = 0,
    concor_location_code: str = "",
) -> str:
    location_text = _clean_text(location)
    location_upper = location_text.upper()
    train_number = _clean_text(train_no)
    departure_text = _clean_text(departure)
    concor_code = _clean_text(concor_location_code).upper()
    if _is_arrived(location_text):
        return "Arrived Birgunj"
    if train_number or departure_text or concor_code == "WGN":
        return "On Rail"
    if _is_port(location_text) or "VIZAG" in location_upper or "VISHAKAPATNAM" in location_upper:
        return "At Port"
    if location_text:
        return "On Rail"
    if delay_days > 5:
        return "At Port"
    return "Hi Seas"


def _has_reached_birgunj(shipment: Shipment) -> bool:
    stored_movement = _normalize_existing_movement(shipment.movement_category or "")
    rail_status = _normalize_existing_movement(shipment.rail_status or "")
    latest_location = _clean_text(shipment.latest_location)
    shipment_status = _clean_text(shipment.shipment_status).lower()

    if (
        stored_movement == "Arrived Birgunj"
        or rail_status == "Arrived Birgunj"
        or _is_arrived(latest_location)
    ):
        return True

    # Finished shipment cycles should preserve the destination milestone even if
    # the empty container later moves away from Birgunj in downstream live feeds.
    return shipment_status in {"completed", "archived"}


def _effective_shipment_movement(shipment: Shipment) -> str:
    if _has_reached_birgunj(shipment):
        return "Arrived Birgunj"

    return _normalize_existing_movement(
        _movement_category(
            shipment.latest_location,
            shipment.train_no,
            shipment.departure,
            shipment.delay_days,
            getattr(shipment, "concor_location_code", ""),
        )
    )


def _effective_shipment_location(shipment: Shipment) -> str:
    tracking_source = _clean_text(getattr(shipment, "tracking_source", "")).lower()
    birgunj_arrival_date = _clean_text(getattr(shipment, "birgunj_arrival_date", ""))
    latest_location = _clean_text(getattr(shipment, "latest_location", ""))
    if "pristine" in tracking_source and birgunj_arrival_date:
        return "ICD BIRGANJ, Samastipur"
    if _has_reached_birgunj(shipment) and _is_arrived(latest_location):
        return "ICD BIRGANJ, Samastipur"
    return latest_location


def _shipment_needs_action(shipment: Shipment) -> bool:
    tracking_source = _clean_text(getattr(shipment, "tracking_source", "")).lower()
    booking_date = _parse_date(_clean_text(getattr(shipment, "pristine_booking_date", "")))
    if "pristine" not in tracking_source or not booking_date:
        return False
    if not _has_reached_birgunj(shipment):
        return False
    arrival_date = _parse_date(_clean_text(getattr(shipment, "birgunj_arrival_date", "")))
    if arrival_date is None and _effective_shipment_movement(shipment) == "Arrived Birgunj":
        arrival_date = _parse_date(_clean_text(getattr(shipment, "latest_time", "")))
    if arrival_date is None:
        return False
    return booking_date > arrival_date


def _tracking_source_labels(tracking_source: str) -> list[str]:
    labels: list[str] = []
    for source in re.split(r"[,+]", _clean_text(tracking_source).lower()):
        source = source.strip()
        if not source:
            continue
        if source == "ldb":
            label = "LDB"
        elif source == "concor":
            label = "CONCOR"
        elif source == "pristine":
            label = "Pristine"
        else:
            label = source.capitalize()
        if label not in labels:
            labels.append(label)
    return labels


def _join_human_labels(labels: list[str]) -> str:
    clean_labels = [label for label in labels if _clean_text(label)]
    if not clean_labels:
        return ""
    if len(clean_labels) == 1:
        return clean_labels[0]
    if len(clean_labels) == 2:
        return f"{clean_labels[0]} and {clean_labels[1]}"
    return f"{', '.join(clean_labels[:-1])}, and {clean_labels[-1]}"


def _movement_resolution_summary(
    shipment: Shipment,
    movement_category: str,
    effective_location: str,
) -> str:
    stored_movement = _normalize_existing_movement(shipment.movement_category or "")
    rail_status = _normalize_existing_movement(shipment.rail_status or "")
    latest_location = _clean_text(shipment.latest_location)
    shipment_status = _clean_text(shipment.shipment_status).lower()
    train_no = _clean_text(shipment.train_no)
    departure = _clean_text(shipment.departure)
    port_arrival_date = _clean_text(shipment.port_arrival_date)
    birgunj_arrival_date = _clean_text(getattr(shipment, "birgunj_arrival_date", ""))
    concor_code = _clean_text(getattr(shipment, "concor_location_code", "")).upper()

    if movement_category == "Arrived Birgunj":
        if birgunj_arrival_date:
            return "Marked as Arrived Birgunj because a Birgunj arrival milestone is saved for this shipment."
        if _is_arrived(latest_location):
            return "Marked as Arrived Birgunj because the latest live location already points to Birgunj."
        if stored_movement == "Arrived Birgunj" or rail_status == "Arrived Birgunj":
            return "Marked as Arrived Birgunj because the saved live movement already confirms Birgunj."
        if shipment_status in {"completed", "archived"}:
            return "Kept at Arrived Birgunj because this shipment cycle is already closed."
        return "Marked as Arrived Birgunj because the freshest saved source signals now point to Birgunj."

    if movement_category == "On Rail":
        if train_no or departure or concor_code == "WGN":
            return "Marked as On Rail because active inland rail movement signals are saved on this shipment."
        if _clean_text(effective_location):
            return "Marked as On Rail because there is a live inland location, but Birgunj arrival is not yet confirmed."
        return "Marked as On Rail because the latest saved tracking fields show the shipment is already moving inland."

    if movement_category == "At Port":
        if port_arrival_date:
            return "Marked as At Port because a port-arrival milestone is saved and Birgunj arrival is not confirmed yet."
        if _is_port(latest_location):
            return "Marked as At Port because the latest saved location is still at the port side."
        return "Marked as At Port because the latest saved movement has not progressed inland yet."

    return "Kept at Hi Seas because no fresh port-side or inland-arrival confirmation is saved on this shipment yet."


def _movement_source_priority_summary(shipment: Shipment) -> str:
    source_labels = _tracking_source_labels(getattr(shipment, "tracking_source", ""))
    lower_sources = {label.lower() for label in source_labels}
    has_live_rail_source = any(label in lower_sources for label in {"ldb", "concor"})
    has_pristine = "pristine" in lower_sources

    if has_live_rail_source and has_pristine:
        return "LDB and CONCOR are checked first. Pristine is only used as a fallback when the live rail feeds are missing, stale, or still lagging behind the latest arrival."
    if has_live_rail_source:
        return "The current movement is being driven by the live rail sources saved on this shipment."
    if has_pristine:
        return "Pristine is providing the current saved signal because fresher live rail confirmation is not available on this shipment."
    return "This shipment is currently relying on the most recently saved portal fields."


def _movement_diagnostic_evidence(
    shipment: Shipment,
    effective_location: str,
    movement_category: str,
    action_required: bool,
) -> list[dict[str, str]]:
    evidence: list[dict[str, str]] = []
    source_text = _join_human_labels(_tracking_source_labels(getattr(shipment, "tracking_source", "")))

    def add(label: str, value: Any) -> None:
        text_value = _clean_text(value)
        if text_value:
            evidence.append({"label": label, "value": text_value})

    add("Resolved movement", movement_category)
    add("Source mix", source_text or "Manual entry")
    add("Saved latest location", effective_location)
    add("Saved latest activity", shipment.latest_time)
    add("Port arrival", shipment.port_arrival_date)
    add("Birgunj arrival", getattr(shipment, "birgunj_arrival_date", ""))
    add("Train number", shipment.train_no)
    add("Departure", shipment.departure)
    add("Wagon loaded", getattr(shipment, "wagon_loaded_date", ""))
    add("Pristine booking", getattr(shipment, "pristine_booking_date", ""))
    add("Latest check status", shipment.last_refresh_status)
    if action_required:
        add("Action note", "Pristine booking date is after Birgunj arrival")
    add("Latest check note", shipment.last_refresh_error)
    return evidence


def _movement_diagnostics(
    shipment: Shipment,
    movement_category: str,
    effective_location: str,
    action_required: bool,
) -> dict[str, Any]:
    existing_diagnostics = getattr(shipment, "_movement_diagnostics", None)
    if isinstance(existing_diagnostics, dict) and existing_diagnostics:
        diagnostics = dict(existing_diagnostics)
        diagnostics.setdefault(
            "resolution_summary",
            _movement_resolution_summary(shipment, movement_category, effective_location),
        )
        diagnostics.setdefault("source_priority_summary", _movement_source_priority_summary(shipment))
        diagnostics.setdefault("group_scope_summary", "")
        diagnostics.setdefault(
            "evidence",
            _movement_diagnostic_evidence(
                shipment,
                effective_location,
                movement_category,
                action_required,
            ),
        )
        return diagnostics
    return {
        "resolution_summary": _movement_resolution_summary(shipment, movement_category, effective_location),
        "source_priority_summary": _movement_source_priority_summary(shipment),
        "group_scope_summary": "",
        "evidence": _movement_diagnostic_evidence(
            shipment,
            effective_location,
            movement_category,
            action_required,
        ),
    }


def _movement_since_date(
    movement_category: str,
    latest_time: str,
    port_arrival_date: str,
    birgunj_arrival_date: str,
    departure: str,
    wagon_loaded_date: str = "",
) -> str:
    movement = _normalize_existing_movement(movement_category or "")
    if movement == "Arrived Birgunj":
        return _clean_text(birgunj_arrival_date) or _clean_text(latest_time)
    if movement == "On Rail":
        return _clean_text(departure) or _clean_text(wagon_loaded_date) or _clean_text(latest_time)
    if movement == "At Port":
        return _clean_text(port_arrival_date) or _clean_text(latest_time)
    return _clean_text(latest_time)


def _extract_json_object_by_key(json_str: str, key_name: str) -> dict[str, Any] | None:
    try:
        key_pattern = f'"{key_name}"'
        key_pos = json_str.find(key_pattern)
        if key_pos == -1:
            return None
        brace_start = json_str.find("{", key_pos)
        if brace_start == -1:
            return None
        brace_count = 0
        in_string = False
        escape_next = False
        for index in range(brace_start, len(json_str)):
            char = json_str[index]
            if escape_next:
                escape_next = False
                continue
            if char == "\\":
                escape_next = True
                continue
            if char == '"':
                in_string = not in_string
                continue
            if in_string:
                continue
            if char == "{":
                brace_count += 1
            elif char == "}":
                brace_count -= 1
                if brace_count == 0:
                    return json.loads(json_str[brace_start : index + 1])
    except (json.JSONDecodeError, TypeError, ValueError):
        return None
    return None


def _json_value(obj: dict[str, Any] | None, key: str, default: str = "") -> str:
    if not isinstance(obj, dict):
        return default
    value = obj.get(key, default)
    return default if value is None else _clean_text(value)


def _format_to_dd_mm_yyyy(date_str: str) -> str:
    if not date_str:
        return ""
    if " " in date_str:
        date_str = date_str.split(" ")[0]
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(date_str, fmt).strftime("%d-%m-%Y")
        except ValueError:
            continue
    return date_str


def _normalize_manual_date(value: Any) -> str:
    text_value = _clean_text(value)
    if not text_value:
        return ""
    return _format_to_dd_mm_yyyy(text_value)


def _extract_concor_wagon_signal(details_text: str, last_reported_text: str = "") -> tuple[str, str]:
    combined_text = " ".join(
        part for part in (_clean_text(details_text), _clean_text(last_reported_text)) if part
    )
    if not combined_text:
        return "", ""

    code_match = re.search(r"\bWGN\b", combined_text, re.IGNORECASE)
    if not code_match:
        return "", ""

    date_match = re.search(r"since\s*\(?\s*(\d{2}/\d{2}/\d{4})", combined_text, re.IGNORECASE)
    wagon_loaded_date = _format_to_dd_mm_yyyy(date_match.group(1)) if date_match else ""
    return "WGN", wagon_loaded_date


def _extract_concor_last_reported_station(last_reported_text: str) -> str:
    text_value = _clean_text(last_reported_text)
    if not text_value:
        return ""

    # Keep the human station text, but strip trailing date/status fragments that
    # are useful for freshness checks and wagon signals, not for display.
    station_text = re.sub(r"\bsince\b.*$", "", text_value, flags=re.IGNORECASE)
    station_text = re.sub(r"\b\d{2}/\d{2}/\d{4}\b", "", station_text)
    station_text = re.sub(r"[(),;]+", " ", station_text)
    station_text = re.sub(r"\s{2,}", " ", station_text).strip(" -:/")

    if station_text.upper() in {"", "WGN", "NA", "N/A", "NIL"}:
        return ""
    return station_text


def _normalize_document_status(value: Any) -> str:
    text_value = _clean_text(value).lower()
    if text_value == "copy":
        return "Copy"
    if text_value == "original":
        return "Original"
    return ""


def _format_ldb_date(iso_text: str) -> str:
    if not iso_text:
        return ""
    try:
        if "T" in iso_text:
            iso_text = iso_text.split("T")[0]
        return datetime.fromisoformat(iso_text).strftime("%d-%m-%Y")
    except ValueError:
        return _format_to_dd_mm_yyyy(iso_text)


def _parse_date(value: str) -> datetime | None:
    text_value = _clean_text(value)
    if not text_value:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text_value, fmt)
        except ValueError:
            continue
    return None


def _parse_ldb_timestamp(value: str) -> datetime | None:
    text_value = _clean_text(value)
    if not text_value:
        return None
    try:
        return datetime.fromisoformat(text_value.replace("Z", "+00:00"))
    except ValueError:
        return None


def _days_between_dates(left: datetime | None, right: datetime | None) -> int | None:
    if left is None or right is None:
        return None
    return abs((left.date() - right.date()).days)


def _freshest_tracking_date(payload: dict[str, Any] | None, keys: list[str]) -> datetime | None:
    freshest: datetime | None = None
    for key in keys:
        parsed = _parse_date(_clean_text((payload or {}).get(key, "")))
        if parsed is None:
            continue
        if freshest is None or parsed > freshest:
            freshest = parsed
    return freshest


def _has_concor_live_signal(concor_data: dict[str, Any] | None) -> bool:
    payload = concor_data or {}
    return any(
        _clean_text(payload.get(field, ""))
        for field in (
            "train_no",
            "departure",
            "wagon_loaded_date",
            "concor_location_code",
            "last_reported_station",
            "last_reported_date",
        )
    )


def _has_ldb_live_signal(ldb_data: dict[str, Any] | None) -> bool:
    payload = ldb_data or {}
    return any(
        _clean_text(payload.get(field, ""))
        for field in ("latest_location", "latest_time", "port_arrival_date", "birgunj_arrival_date", "rail_status")
    )


def _is_tracking_date_stale(candidate: datetime | None) -> bool:
    if candidate is None:
        return False
    current_date = datetime.now(APP_TIMEZONE).replace(tzinfo=None)
    days_from_now = _days_between_dates(current_date, candidate)
    return bool(days_from_now is not None and days_from_now >= STALE_PORTAL_RECORD_DAYS)


def _should_ignore_stale_ldb_data(
    ldb_data: dict[str, Any] | None,
    concor_data: dict[str, Any] | None,
    pristine_data: dict[str, Any] | None,
) -> bool:
    payload = ldb_data or {}
    if not _has_ldb_live_signal(payload):
        return False

    freshest_ldb_date = _freshest_tracking_date(payload, ["latest_time", "port_arrival_date", "birgunj_arrival_date"])
    if freshest_ldb_date is None:
        return False

    freshest_concor_date = _freshest_tracking_date(concor_data or {}, ["departure", "wagon_loaded_date", "last_reported_date"])
    freshest_pristine_date = _freshest_tracking_date(
        pristine_data or {},
        ["arrival_date", "booking_date", "empty_date", "rake_departure_date"],
    )

    if _is_tracking_date_stale(freshest_ldb_date) and not any(
        not _is_tracking_date_stale(reference_date)
        for reference_date in (freshest_concor_date, freshest_pristine_date)
        if reference_date is not None
    ):
        return True

    days_from_concor = _days_between_dates(freshest_ldb_date, freshest_concor_date)
    if days_from_concor is not None and days_from_concor >= STALE_PORTAL_RECORD_DAYS:
        return True

    days_from_pristine = _days_between_dates(freshest_ldb_date, freshest_pristine_date)
    if days_from_pristine is not None and days_from_pristine >= STALE_PORTAL_RECORD_DAYS:
        return True

    return False


def _should_ignore_stale_concor_data(
    concor_data: dict[str, Any] | None,
    ldb_data: dict[str, Any] | None,
    pristine_data: dict[str, Any] | None,
) -> bool:
    payload = concor_data or {}
    if not _has_concor_live_signal(payload):
        return False

    freshest_concor_date = _freshest_tracking_date(payload, ["departure", "wagon_loaded_date", "last_reported_date"])
    if freshest_concor_date is None:
        return False

    freshest_ldb_date = _freshest_tracking_date(ldb_data or {}, ["latest_time", "port_arrival_date", "birgunj_arrival_date"])
    freshest_pristine_date = _freshest_tracking_date(
        pristine_data or {},
        ["arrival_date", "booking_date", "empty_date", "rake_departure_date"],
    )

    if _is_tracking_date_stale(freshest_concor_date) and not any(
        not _is_tracking_date_stale(reference_date)
        for reference_date in (freshest_ldb_date, freshest_pristine_date)
        if reference_date is not None
    ):
        return True

    days_from_ldb = _days_between_dates(freshest_concor_date, freshest_ldb_date)
    if days_from_ldb is not None and days_from_ldb >= STALE_PORTAL_RECORD_DAYS:
        return True

    days_from_pristine = _days_between_dates(freshest_concor_date, freshest_pristine_date)
    if days_from_pristine is not None and days_from_pristine >= STALE_PORTAL_RECORD_DAYS:
        return True

    return False


def _should_ignore_stale_pristine_data(
    pristine_data: dict[str, Any] | None,
    ldb_data: dict[str, Any] | None,
    concor_data: dict[str, Any] | None,
) -> bool:
    payload = pristine_data or {}
    if not payload.get("arrival_date"):
        return False

    freshest_pristine_date = _freshest_tracking_date(
        payload,
        ["arrival_date", "booking_date", "empty_date", "rake_departure_date"],
    )
    if freshest_pristine_date is None:
        return False

    ldb_latest_date = _parse_date(_clean_text((ldb_data or {}).get("latest_time", "")))
    ldb_birgunj_date = _parse_date(_clean_text((ldb_data or {}).get("birgunj_arrival_date", "")))
    ldb_port_date = _parse_date(_clean_text((ldb_data or {}).get("port_arrival_date", "")))
    current_date = datetime.now(APP_TIMEZONE).replace(tzinfo=None)

    days_from_now = _days_between_dates(current_date, freshest_pristine_date) or 0
    days_from_ldb_latest = _days_between_dates(ldb_latest_date, freshest_pristine_date)
    days_from_ldb_port = _days_between_dates(ldb_port_date, freshest_pristine_date)

    has_current_birgunj_confirmation = ldb_birgunj_date is not None
    has_other_live_signal = bool(ldb_latest_date or ldb_port_date or _has_concor_live_signal(concor_data))

    if has_current_birgunj_confirmation:
        return False

    if days_from_now >= STALE_PORTAL_RECORD_DAYS and not has_other_live_signal:
        return True

    if days_from_ldb_latest is not None and days_from_ldb_latest >= STALE_PORTAL_RECORD_DAYS:
        return True

    if days_from_ldb_port is not None and days_from_ldb_port >= STALE_PORTAL_RECORD_DAYS:
        return True

    return False


def _freshest_live_tracking_date(
    ldb_data: dict[str, Any] | None,
    concor_data: dict[str, Any] | None,
) -> datetime | None:
    candidates = [
        _freshest_tracking_date(ldb_data or {}, ["latest_time", "port_arrival_date", "birgunj_arrival_date"]),
        _freshest_tracking_date(concor_data or {}, ["departure", "wagon_loaded_date", "last_reported_date"]),
    ]
    parsed_dates = [candidate for candidate in candidates if candidate is not None]
    if not parsed_dates:
        return None
    return max(parsed_dates)


def _should_use_pristine_arrival_override(
    pristine_data: dict[str, Any] | None,
    ldb_data: dict[str, Any] | None,
    concor_data: dict[str, Any] | None,
) -> bool:
    pristine_candidates = extract_pristine_candidates(pristine_data)
    ldb_candidates = extract_ldb_candidates(ldb_data)
    concor_candidates = extract_concor_candidates(concor_data)

    pristine_birgunj = next(
        (candidate for candidate in pristine_candidates if candidate.milestone == "birgunj_arrival"),
        None,
    )
    live_birgunj = next(
        (
            candidate
            for candidate in [*ldb_candidates, *concor_candidates]
            if candidate.milestone == "birgunj_arrival"
        ),
        None,
    )
    live_inland = next(
        (
            candidate
            for candidate in [*ldb_candidates, *concor_candidates]
            if candidate.milestone == "inland_movement"
        ),
        None,
    )
    accepted_port = next(
        (
            candidate
            for candidate in [*ldb_candidates, *concor_candidates]
            if candidate.milestone == "port_arrival"
        ),
        None,
    )

    for candidate in (pristine_birgunj, live_birgunj, live_inland):
        if candidate is not None and candidate.date is not None:
            candidate.stale_by_age = _is_tracking_date_stale(candidate.date)
            candidate.same_cycle = not candidate.stale_by_age
            if (
                candidate is pristine_birgunj
                and accepted_port is not None
                and accepted_port.date is not None
                and candidate.date < accepted_port.date
            ):
                candidate.same_cycle = False
            if candidate.date and candidate.milestone == "birgunj_arrival" and live_inland and live_inland.date:
                candidate.same_cycle = candidate.same_cycle and candidate.date >= live_inland.date - timedelta(days=MAX_SAME_CYCLE_BACKWARD_DRIFT_DAYS)

    return pristine_birgunj_can_override_lagging_live_feeds(pristine_birgunj, live_birgunj, live_inland)


def _date_sort_value(value: str) -> float:
    parsed = _parse_date(value)
    return parsed.timestamp() if parsed else 0.0


def _earliest_non_empty_date(values: list[str]) -> str:
    dated = [(_date_sort_value(value), value) for value in values if _clean_text(value)]
    dated = [item for item in dated if item[0] > 0]
    if not dated:
        return _first_non_empty(values)
    return min(dated, key=lambda item: item[0])[1]


def _entry_event_text(entry: dict[str, Any]) -> str:
    return _json_value(entry, "eventName").upper()


def _entry_location_text(entry: dict[str, Any]) -> str:
    return _json_value(entry, "currentLocation")


def _is_origin_port_event(entry: dict[str, Any]) -> bool:
    entry_location = _entry_location_text(entry)
    entry_event = _entry_event_text(entry)
    if not entry_location or _is_arrived(entry_location):
        return False
    if not _is_port(entry_location):
        return False
    return any(keyword in entry_event for keyword in ("PORT IN", "PORT OUT", "ICD IN", "CFS IN", "CFS OUT"))


def _is_birgunj_event(entry: dict[str, Any]) -> bool:
    return _is_arrived(_entry_location_text(entry))


def _derive_ldb_milestones(last_event: dict[str, Any], track_log: list[dict[str, Any]] | None) -> dict[str, str]:
    latest_date = _format_ldb_date(_json_value(last_event, "timestampTimezone"))
    if not isinstance(track_log, list) or not track_log:
        return {
            "latest_location": _json_value(last_event, "currentLocation"),
            "latest_time": latest_date,
            "port_arrival_date": "",
            "birgunj_arrival_date": latest_date if _is_arrived(_json_value(last_event, "currentLocation")) else "",
        }

    ordered_entries = []
    for entry in track_log:
        if not isinstance(entry, dict):
            continue
        parsed_timestamp = _parse_ldb_timestamp(_json_value(entry, "timestampTimezone"))
        if not parsed_timestamp:
            continue
        ordered_entries.append((parsed_timestamp, entry))

    if not ordered_entries:
        return {
            "latest_location": _json_value(last_event, "currentLocation"),
            "latest_time": latest_date,
            "port_arrival_date": "",
            "birgunj_arrival_date": latest_date if _is_arrived(_json_value(last_event, "currentLocation")) else "",
        }

    ordered_entries.sort(key=lambda item: item[0])
    entries = [entry for _, entry in ordered_entries]

    birgunj_indices = [index for index, entry in enumerate(entries) if _is_birgunj_event(entry)]
    latest_birgunj_index = birgunj_indices[-1] if birgunj_indices else -1

    boundary_index = latest_birgunj_index
    if latest_birgunj_index >= 0:
        last_origin_before_latest_birgunj = max(
            (index for index, entry in enumerate(entries[:latest_birgunj_index]) if _is_origin_port_event(entry)),
            default=-1,
        )
        if last_origin_before_latest_birgunj >= 0:
            boundary_index = max(
                (index for index in birgunj_indices if index < last_origin_before_latest_birgunj),
                default=-1,
            )

    relevant_entries = entries[boundary_index + 1 :] if boundary_index >= 0 else entries

    port_in_dates: list[str] = []
    port_out_dates: list[str] = []
    origin_icd_in_dates: list[str] = []
    birgunj_dates: list[str] = []

    for entry in relevant_entries:
        entry_location = _entry_location_text(entry)
        entry_event = _entry_event_text(entry)
        formatted_entry_date = _format_ldb_date(_json_value(entry, "timestampTimezone"))
        if not formatted_entry_date:
            continue
        if _is_arrived(entry_location):
            birgunj_dates.append(formatted_entry_date)
            continue
        if not _is_port(entry_location):
            continue
        if "PORT IN" in entry_event:
            port_in_dates.append(formatted_entry_date)
        elif "PORT OUT" in entry_event:
            port_out_dates.append(formatted_entry_date)
        elif "ICD IN" in entry_event:
            origin_icd_in_dates.append(formatted_entry_date)

    return {
        "latest_location": _json_value(last_event, "currentLocation"),
        "latest_time": latest_date,
        "port_arrival_date": _earliest_non_empty_date(port_in_dates)
        or _earliest_non_empty_date(port_out_dates)
        or _earliest_non_empty_date(origin_icd_in_dates),
        "birgunj_arrival_date": _earliest_non_empty_date(birgunj_dates),
    }


def _classify_ldb_rail_status(location: str, event: str) -> str:
    text_value = f"{location} {event}".upper()
    if "BIRGANJ" in text_value or "BIRGUNJ" in text_value:
        return "Arrived Birgunj"
    if "DANGOAPOSI" in text_value or "STATION CROSSED" in text_value or "RAIL" in text_value:
        return "On Rail"
    if _is_port(text_value):
        return "At Port"
    return "On Rail" if _clean_text(location) else "Hi Seas"


def _compute_delay_days(latest_date: str) -> float:
    parsed = _parse_date(latest_date)
    if not parsed:
        return 0
    return round((datetime.now() - parsed).total_seconds() / 86400, 2)


def _get_cached_data(container_number: str) -> dict[str, Any] | None:
    if not CACHE_FILE.exists():
        return None
    try:
        cache = json.loads(CACHE_FILE.read_text(encoding="utf-8"))
        cached = cache.get(container_number)
        if not cached:
            return None
        cached_time = datetime.fromisoformat(cached.get("cached_at", "2000-01-01T00:00:00"))
        if datetime.now() - cached_time < timedelta(minutes=CACHE_DURATION_MINUTES):
            cached_data = cached.get("data")
            if isinstance(cached_data, dict) and {"ldb", "concor", "pristine"}.issubset(cached_data.keys()):
                return cached_data
    except (OSError, json.JSONDecodeError, TypeError, ValueError) as exc:
        logger.warning("Unable to read tracking cache for %s: %s", container_number, exc)
        return None
    return None


def _save_to_cache(container_number: str, data: dict[str, Any]) -> None:
    try:
        cache = json.loads(CACHE_FILE.read_text(encoding="utf-8")) if CACHE_FILE.exists() else {}
        cache[container_number] = {"data": data, "cached_at": _now_iso_for_cache()}
        CACHE_FILE.write_text(json.dumps(cache, indent=2), encoding="utf-8")
    except (OSError, json.JSONDecodeError, TypeError, ValueError) as exc:
        logger.warning("Unable to save tracking cache for %s: %s", container_number, exc)
        return


def _fetch_tracking_sources(container_number: str, use_cache: bool = True) -> dict[str, Any]:
    container_number = _clean_container(container_number)
    if not container_number:
        return {
            "ldb": {},
            "concor": {},
            "pristine": {},
            "status": "error",
            "error": "Missing container number",
            "cached": False,
            "has_data": False,
        }

    if use_cache:
        cached_sources = _get_cached_data(container_number)
        if cached_sources:
            has_data = bool(
                (cached_sources.get("ldb") or {})
                or (cached_sources.get("concor") or {})
                or (cached_sources.get("pristine") or {})
            )
            return {
                "ldb": cached_sources.get("ldb") or {},
                "concor": cached_sources.get("concor") or {},
                "pristine": cached_sources.get("pristine") or {},
                "status": "success-cached" if has_data else "no_data",
                "error": "" if has_data else "No data from any API",
                "cached": True,
                "has_data": has_data,
            }

    with ThreadPoolExecutor(max_workers=3) as executor:
        future_ldb = executor.submit(_fetch_ldb, container_number)
        future_concor = executor.submit(_fetch_concor, container_number)
        future_pristine = executor.submit(_fetch_pristine_arrival, container_number)
        ldb_data = future_ldb.result() or {}
        concor_data = future_concor.result() or {}
        pristine_data = future_pristine.result() or {}

    if _should_ignore_stale_ldb_data(ldb_data, concor_data, pristine_data):
        ldb_data = {}

    if _should_ignore_stale_concor_data(concor_data, ldb_data, pristine_data):
        concor_data = {}

    if _should_ignore_stale_pristine_data(pristine_data, ldb_data, concor_data):
        pristine_data = {}

    sources = {"ldb": ldb_data, "concor": concor_data, "pristine": pristine_data}
    has_data = bool(ldb_data or concor_data or pristine_data)
    error = (
        ""
        if _clean_text(concor_data.get("train_no", "")) or _clean_text(concor_data.get("concor_location_code", "")) == "WGN" or not concor_data
        else "CONCOR returned but no train number"
    )
    if not has_data:
        error = "No data from any API"

    if has_data:
        _save_to_cache(container_number, sources)

    return {
        "ldb": ldb_data,
        "concor": concor_data,
        "pristine": pristine_data,
        "status": "success" if has_data else "no_data",
        "error": error,
        "cached": False,
        "has_data": has_data,
    }


def _fetch_ldb(container_number: str) -> dict[str, Any] | None:
    try:
        response = requests.get(
            f"{LDB_API_URL}?cntrNo={quote(container_number)}&searchType=39",
            timeout=(4, TRACKING_SOURCE_TIMEOUT_SECONDS),
            headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36", "Accept": "application/json, text/plain, */*", "Referer": f"https://www.ldb.co.in/ldb/containersearch/39/{container_number}"},
        )
        if response.status_code != 200 or not response.text:
            return None
        payload_json = None
        try:
            payload_json = response.json()
        except ValueError:
            payload_json = None

        last_event = _extract_json_object_by_key(response.text, "lastEvent")
        track_log = None
        if not last_event:
            data = payload_json
            if isinstance(data, dict):
                last_event = data.get("lastEvent") or data.get("data") or data.get("result")
                if isinstance(last_event, list) and last_event:
                    last_event = last_event[0]
        if isinstance(payload_json, dict):
            object_payload = payload_json.get("object") if isinstance(payload_json.get("object"), dict) else payload_json
            candidate_track_log = object_payload.get("trackLog") if isinstance(object_payload, dict) else None
            if isinstance(candidate_track_log, list):
                track_log = candidate_track_log
        if not isinstance(last_event, dict):
            return None
        milestones = _derive_ldb_milestones(last_event, track_log)
        location = milestones["latest_location"]
        event = _json_value(last_event, "eventName")
        latest_date = milestones["latest_time"]
        return {
            "latest_location": location,
            "latest_time": latest_date,
            "port_arrival_date": milestones["port_arrival_date"],
            "birgunj_arrival_date": milestones["birgunj_arrival_date"],
            "rail_status": _classify_ldb_rail_status(location, event),
            "delay_days": _compute_delay_days(latest_date),
        }
    except requests.RequestException as exc:
        logger.warning("LDB fetch failed for %s: %s", container_number, exc)
        return None
    except (TypeError, ValueError) as exc:
        logger.warning("LDB payload parse failed for %s: %s", container_number, exc)
        return None


def _fetch_concor(container_number: str) -> dict[str, Any] | None:
    try:
        response = requests.post(
            CONCOR_API_URL,
            json={"containerNo": [container_number]},
            timeout=(4, TRACKING_SOURCE_TIMEOUT_SECONDS),
            headers={"Content-Type": "application/json", "Accept": "application/json, text/plain, */*", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36", "Origin": "https://www.concorindia.co.in", "Referer": "https://www.concorindia.co.in/track-n-trace?lang=en"},
        )
        if response.status_code != 200:
            return None
        payload = response.json()
        if not isinstance(payload, dict) or payload.get("statusCode") == 404 or payload.get("status") == "error":
            return None
        actual_data = payload.get("data")
        container_track = None
        if isinstance(actual_data, dict):
            container_info = actual_data.get(container_number)
            if isinstance(container_info, dict):
                container_track = container_info.get("containerTrack") or container_info
            if not isinstance(container_track, dict):
                for value in actual_data.values():
                    if isinstance(value, dict):
                        if "containerTrack" in value and isinstance(value["containerTrack"], dict):
                            container_track = value["containerTrack"]
                            break
                        if "TRAIN_NUMBER" in value:
                            container_track = value
                            break
            if not isinstance(container_track, dict) and "TRAIN_NUMBER" in actual_data:
                container_track = actual_data
        if not isinstance(container_track, dict):
            return None
        departure = _format_to_dd_mm_yyyy(_json_value(container_track, "DEPARTURE_DATE_&_TIME"))
        last_reported_raw = _json_value(container_track, "LAST_REPORTED_STATION")
        details_text = _json_value(container_track, "DETAILS") or _json_value(container_track, "details")
        concor_location_code, wagon_loaded_date = _extract_concor_wagon_signal(details_text, last_reported_raw)
        last_reported_station = _extract_concor_last_reported_station(last_reported_raw)
        last_reported_date = ""
        date_match = re.search(r"(\d{2}/\d{2}/\d{4})", last_reported_raw)
        if date_match:
            last_reported_date = _format_to_dd_mm_yyyy(date_match.group(1))
        return {
            "train_no": _json_value(container_track, "TRAIN_NUMBER"),
            "wagon_no": _json_value(container_track, "WAGON_NUMBER"),
            "train_origin": _json_value(container_track, "TRAIN_ORIGNATING_STATION"),
            "train_destination": _json_value(container_track, "TRAIN_DESTINATION_STATION"),
            "departure": departure,
            "last_reported_station": last_reported_station,
            "last_reported_date": last_reported_date,
            "shipping_line": "",
            "concor_location_code": concor_location_code,
            "wagon_loaded_date": wagon_loaded_date,
        }
    except requests.RequestException as exc:
        logger.warning("CONCOR fetch failed for %s: %s", container_number, exc)
        return None
    except (TypeError, ValueError) as exc:
        logger.warning("CONCOR payload parse failed for %s: %s", container_number, exc)
        return None


def _extract_pristine_tracking_field(html_text: str, label: str) -> str:
    match = re.search(
        rf"<span><b>{re.escape(label)}\s*:?\s*</b></span>\s*<span>(.*?)</span>",
        html_text,
        re.IGNORECASE | re.DOTALL,
    )
    if not match:
        return ""
    raw_value = re.sub(r"<.*?>", "", match.group(1)).strip()
    return unescape(raw_value)


def _extract_pristine_nonce(html_text: str) -> str:
    match = re.search(
        r'<input[^>]*name="_wpnonce_phoen_tracking"[^>]*value="([^"]+)"',
        html_text,
        re.IGNORECASE,
    )
    if not match:
        match = re.search(
            r'<input[^>]*value="([^"]+)"[^>]*name="_wpnonce_phoen_tracking"',
            html_text,
            re.IGNORECASE,
        )
    return match.group(1) if match else ""


def _fetch_pristine_arrival(container_number: str) -> dict[str, Any] | None:
    try:
        session = requests.Session()
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/135.0.0.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
            "Upgrade-Insecure-Requests": "1",
            "Referer": PRISTINE_TRACKING_URL,
        }
        page = session.get(
            PRISTINE_TRACKING_URL,
            headers=headers,
            timeout=(4, TRACKING_SOURCE_TIMEOUT_SECONDS),
        )
        if page.status_code != 200 or not page.text:
            return None
        nonce = _extract_pristine_nonce(page.text)
        response = session.post(
            PRISTINE_TRACKING_URL,
            headers=headers,
            data={
                "_wpnonce_phoen_tracking": nonce,
                "tracking_id": container_number,
                "tracking_button": "Track",
            },
            timeout=(4, TRACKING_SOURCE_TIMEOUT_SECONDS),
        )
        if response.status_code != 200 or not response.text:
            return None
        html_text = response.text
        if container_number.upper() not in html_text.upper():
            return None
        arrival_date = _extract_pristine_tracking_field(html_text, "Arrival Date")
        if not arrival_date:
            return None
        return {
            "arrival_date": _format_to_dd_mm_yyyy(arrival_date),
            "booking_date": _format_to_dd_mm_yyyy(_extract_pristine_tracking_field(html_text, "Booking Date")),
            "empty_date": _format_to_dd_mm_yyyy(_extract_pristine_tracking_field(html_text, "Empty Date")),
            "rake_departure_date": _format_to_dd_mm_yyyy(
                _extract_pristine_tracking_field(html_text, "Rake Departure Date")
            ),
            "location": "ICD BIRGANJ, Samastipur",
        }
    except requests.RequestException as exc:
        logger.warning("Pristine fetch failed for %s: %s", container_number, exc)
        return None


def _title_case_word(word: str) -> str:
    if not word:
        return ""
    if "*" in word:
        return word.upper()
    stripped = word.replace(".", "")
    if stripped.upper() in BUSINESS_SUFFIXES:
        return stripped.upper() + ("." if word.endswith(".") else "")
    if len(word) == 1 and word.isalpha():
        return word.upper()
    return word[:1].upper() + word[1:].lower()


def _format_customer_name(value: str) -> str:
    compact = _clean_text(value).replace("_", " ")
    compact = re.sub(r"\s*[--]+\s*", " ", compact)
    compact = re.sub(r"\s+", " ", compact).strip()
    if not compact:
        return ""
    formatted = " ".join(_title_case_word(word) for word in compact.split(" ") if word)
    formatted = re.sub(r"\bPVT\.?\s+LTD\.?\b\.?", "PVT. LTD.", formatted)
    formatted = re.sub(r"\bCO\.?\s+LTD\.?\b\.?", "CO. LTD.", formatted)
    formatted = re.sub(r"\.{2,}", ".", formatted)
    return formatted


def _customer_tokens(value: str) -> list[str]:
    cleaned = re.sub(r"[^A-Z0-9 ]+", " ", _clean_text(value).upper())
    tokens = [token for token in cleaned.split() if token]
    reduced = [token for token in tokens if token not in CUSTOMER_IGNORE_TOKENS]
    return reduced or tokens


def _customer_compact_key(value: str) -> str:
    return "".join(_customer_tokens(value))


def _customer_match_score(left: str, right: str) -> float:
    left_tokens = _customer_tokens(left)
    right_tokens = _customer_tokens(right)
    left_compact = "".join(left_tokens)
    right_compact = "".join(right_tokens)
    if not left_compact or not right_compact:
        return 0.0
    sequence_score = SequenceMatcher(None, left_compact, right_compact).ratio()
    left_set = set(left_tokens)
    right_set = set(right_tokens)
    overlap = len(left_set & right_set) / max(len(left_set | right_set), 1)
    return max(sequence_score, (sequence_score + overlap) / 2)


def _is_customer_match(left: str, right: str) -> bool:
    left_tokens = _customer_tokens(left)
    right_tokens = _customer_tokens(right)
    if not left_tokens or not right_tokens:
        return False
    if len(left_tokens) != len(right_tokens):
        return _customer_match_score(left, right) >= 0.9
    token_scores = [
        SequenceMatcher(None, left_token, right_token).ratio()
        for left_token, right_token in zip(left_tokens, right_tokens)
    ]
    average_token_score = sum(token_scores) / max(len(token_scores), 1)
    return _customer_match_score(left, right) >= 0.88 or average_token_score >= 0.92


def _aliases_from_json(raw_value: str) -> list[str]:
    try:
        parsed = json.loads(raw_value or "[]")
    except (TypeError, json.JSONDecodeError):
        return []
    if not isinstance(parsed, list):
        return []
    return [item for item in parsed if isinstance(item, str) and item.strip()]


def _pick_better_display_name(current: str, candidate: str) -> str:
    current_name = _format_customer_name(current)
    candidate_name = _format_customer_name(candidate)
    if not current_name:
        return candidate_name
    if not candidate_name:
        return current_name
    current_penalty = current_name.count("..") + current_name.count(" .")
    candidate_penalty = candidate_name.count("..") + candidate_name.count(" .")
    current_score = (
        0 if "*" in current_name else 1,
        -current_penalty,
        len(set(_customer_tokens(current_name))),
        len(current_name),
    )
    candidate_score = (
        0 if "*" in candidate_name else 1,
        -candidate_penalty,
        len(set(_customer_tokens(candidate_name))),
        len(candidate_name),
    )
    return candidate_name if candidate_score > current_score else current_name


def _sync_customer_directory(db: Session, raw_name: str) -> str:
    formatted_name = _format_customer_name(raw_name)
    if not formatted_name:
        return ""
    compact_key = _customer_compact_key(formatted_name)
    if not compact_key:
        return formatted_name
    entries = list(
        db.execute(
            select(CustomerDirectory).order_by(
                CustomerDirectory.source_count.desc(),
                CustomerDirectory.updated_at.desc(),
            )
        ).scalars()
    )
    best_entry = None
    best_score = 0.0
    for entry in entries:
        for alias in [entry.display_name, *_aliases_from_json(entry.aliases_json)]:
            score = _customer_match_score(formatted_name, alias)
            if score > best_score:
                best_score = score
                best_entry = entry
    if best_entry and _is_customer_match(formatted_name, best_entry.display_name):
        aliases = _aliases_from_json(best_entry.aliases_json)
        if formatted_name not in aliases and formatted_name != best_entry.display_name:
            aliases.append(formatted_name)
            best_entry.aliases_json = json.dumps(sorted(set(aliases)), ensure_ascii=True)
        best_entry.display_name = _pick_better_display_name(best_entry.display_name, formatted_name)
        best_entry.source_count = (best_entry.source_count or 0) + 1
        return best_entry.display_name
    entry = db.execute(select(CustomerDirectory).where(CustomerDirectory.canonical_key == compact_key)).scalar_one_or_none()
    if entry:
        aliases = _aliases_from_json(entry.aliases_json)
        if formatted_name not in aliases and formatted_name != entry.display_name:
            aliases.append(formatted_name)
            entry.aliases_json = json.dumps(sorted(set(aliases)), ensure_ascii=True)
        entry.display_name = _pick_better_display_name(entry.display_name, formatted_name)
        entry.source_count = (entry.source_count or 0) + 1
        return entry.display_name
    entry = CustomerDirectory(
        canonical_key=compact_key,
        display_name=formatted_name,
        aliases_json=json.dumps([formatted_name], ensure_ascii=True),
        source_count=1,
    )
    db.add(entry)
    db.flush()
    return entry.display_name


def _load_documents_index() -> dict[str, dict[str, Any]]:
    global _documents_index_cache, _documents_index_mtime
    if not BL_DOCUMENTS_INDEX_FILE.exists():
        _documents_index_cache = {}
        _documents_index_mtime = None
        return {}
    try:
        current_mtime = BL_DOCUMENTS_INDEX_FILE.stat().st_mtime
        if _documents_index_cache is not None and _documents_index_mtime == current_mtime:
            return _documents_index_cache
        data = json.loads(BL_DOCUMENTS_INDEX_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, TypeError) as exc:
        logger.warning("Unable to read BL documents index from %s: %s", BL_DOCUMENTS_INDEX_FILE, exc)
        return {}
    normalized = data if isinstance(data, dict) else {}
    _documents_index_cache = normalized
    _documents_index_mtime = current_mtime
    return normalized


def _save_documents_index(data: dict[str, dict[str, Any]]) -> None:
    BL_DOCUMENTS_INDEX_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")
    global _documents_index_cache, _documents_index_mtime
    _documents_index_cache = data
    try:
        _documents_index_mtime = BL_DOCUMENTS_INDEX_FILE.stat().st_mtime
    except OSError as exc:
        logger.warning("Unable to stat BL documents index %s: %s", BL_DOCUMENTS_INDEX_FILE, exc)
        _documents_index_mtime = None


def _r2_bucket_name() -> str:
    return _clean_text(os.getenv("R2_BUCKET_NAME"))


def _r2_endpoint_url() -> str:
    return _clean_text(os.getenv("R2_ENDPOINT_URL"))


def _r2_access_key_id() -> str:
    return _clean_text(os.getenv("R2_ACCESS_KEY_ID"))


def _r2_secret_access_key() -> str:
    return _clean_text(os.getenv("R2_SECRET_ACCESS_KEY"))


def _r2_enabled() -> bool:
    return all(
        [
            _r2_bucket_name(),
            _r2_endpoint_url(),
            _r2_access_key_id(),
            _r2_secret_access_key(),
        ]
    )


def _get_r2_client():
    global _r2_client
    if not _r2_enabled():
        return None
    if _r2_client is None:
        _r2_client = boto3.client(
            "s3",
            endpoint_url=_r2_endpoint_url(),
            aws_access_key_id=_r2_access_key_id(),
            aws_secret_access_key=_r2_secret_access_key(),
            region_name="auto",
        )
    return _r2_client


def _r2_object_key(stored_name: str) -> str:
    return f"bl_documents/{stored_name}"


def _delete_document_object(metadata: dict[str, Any] | None) -> None:
    if not isinstance(metadata, dict):
        return
    storage_backend = _clean_text(metadata.get("storage_backend") or "local").lower()
    if storage_backend == "r2":
        object_key = _clean_text(metadata.get("object_key"))
        client = _get_r2_client()
        bucket_name = _r2_bucket_name()
        if client and object_key:
            try:
                client.delete_object(Bucket=bucket_name, Key=object_key)
            except (BotoCoreError, ClientError) as exc:
                logger.warning("Unable to delete R2 document object %s: %s", object_key, exc)
        return

    target_path = Path(metadata.get("path", ""))
    if target_path.exists() and target_path.is_file():
        target_path.unlink(missing_ok=True)


def _log_audit_event(
    db: Session,
    user_id: int,
    action: str,
    *,
    bl_number: str = "",
    container_number: str = "",
    shipment_status: str = "",
    details: dict[str, Any] | None = None,
) -> None:
    db.add(
        AuditLog(
            user_id=user_id,
            action=_clean_text(action),
            bl_number=_clean_bl(bl_number),
            container_number=_clean_container(container_number),
            shipment_status=_clean_text(shipment_status).lower(),
            details_json=json.dumps(details or {}, default=str),
        )
    )


def _serialize_audit_log(entry: AuditLog) -> dict[str, Any]:
    try:
        details = json.loads(entry.details_json or "{}")
    except (TypeError, json.JSONDecodeError):
        details = {}
    if not isinstance(details, dict):
        details = {"value": details}
    return {
        "id": entry.id,
        "action": entry.action,
        "bl_number": entry.bl_number,
        "container_number": entry.container_number,
        "shipment_status": entry.shipment_status,
        "details": details,
        "created_at": _format_app_timestamp(entry.created_at),
    }


def _documents_for_bl(bl_number: str) -> dict[str, Any]:
    normalized_bl = _normalize_bl_number(bl_number)
    if not normalized_bl:
        return {doc_type: None for doc_type in VALID_DOCUMENT_TYPES}
    existing = _load_documents_index().get(normalized_bl) or {}
    return {
        doc_type: _serialize_document_metadata(normalized_bl, doc_type, existing.get(doc_type))
        for doc_type in VALID_DOCUMENT_TYPES
    }


def _serialize_document_metadata(bl_number: str, document_type: str, metadata: dict[str, Any] | None) -> dict[str, Any] | None:
    if not isinstance(metadata, dict):
        return None
    payload = dict(metadata)
    payload["file_url"] = (
        f"/api/shipments/bl-documents/file?bl_number={quote(bl_number)}&document_type={quote(document_type)}"
    )
    return payload


def _save_document(bl_number: str, document_type: str, upload: UploadFile) -> dict[str, Any]:
    normalized_bl = _normalize_bl_number(bl_number)
    normalized_type = _clean_text(document_type).lower()
    if not normalized_bl:
        raise HTTPException(status_code=400, detail="BL number is required for document upload")
    if normalized_type not in VALID_DOCUMENT_TYPES:
        raise HTTPException(status_code=400, detail="Invalid document_type")
    suffix = Path(upload.filename or "").suffix or ".bin"
    stored_name = f"{normalized_bl}_{normalized_type}_{uuid4().hex}{suffix}"
    content = upload.file.read()
    metadata = {
        "document_type": normalized_type,
        "original_name": upload.filename or stored_name,
        "stored_name": stored_name,
        "content_type": upload.content_type or "application/octet-stream",
        "size": len(content),
        "uploaded_at": _now_datetime(),
    }
    if _r2_enabled():
        object_key = _r2_object_key(stored_name)
        client = _get_r2_client()
        bucket_name = _r2_bucket_name()
        if client is None:
            raise HTTPException(status_code=500, detail="Document storage is not available")
        client.put_object(
            Bucket=bucket_name,
            Key=object_key,
            Body=content,
            ContentType=metadata["content_type"],
        )
        metadata.update(
            {
                "storage_backend": "r2",
                "bucket": bucket_name,
                "object_key": object_key,
            }
        )
    else:
        target_path = BL_DOCUMENTS_DIR / stored_name
        target_path.write_bytes(content)
        metadata.update(
            {
                "storage_backend": "local",
                "path": str(target_path),
            }
        )
    index = _load_documents_index()
    bucket = index.setdefault(normalized_bl, {})
    previous = bucket.get(normalized_type)
    _delete_document_object(previous)
    bucket[normalized_type] = metadata
    _save_documents_index(index)
    return _serialize_document_metadata(normalized_bl, normalized_type, metadata)


def _move_documents_between_bls(previous_bl: str, next_bl: str) -> None:
    previous_normalized = _normalize_bl_number(previous_bl)
    next_normalized = _normalize_bl_number(next_bl)
    if not previous_normalized or not next_normalized or previous_normalized == next_normalized:
        return
    index = _load_documents_index()
    existing = index.get(previous_normalized)
    if not isinstance(existing, dict) or not existing:
        return
    destination = index.setdefault(next_normalized, {})
    for document_type, metadata in existing.items():
        destination[document_type] = metadata
    index.pop(previous_normalized, None)
    _save_documents_index(index)


def _load_location_distance_cache() -> dict[str, dict[str, Any]]:
    if not LOCATION_DISTANCE_CACHE_FILE.exists():
        return {}
    try:
        data = json.loads(LOCATION_DISTANCE_CACHE_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, TypeError) as exc:
        logger.warning(
            "Unable to read location distance cache from %s: %s",
            LOCATION_DISTANCE_CACHE_FILE,
            exc,
        )
        return {}
    return data if isinstance(data, dict) else {}


def _save_location_distance_cache(cache: dict[str, dict[str, Any]]) -> None:
    LOCATION_DISTANCE_CACHE_FILE.write_text(json.dumps(cache, indent=2), encoding="utf-8")


def _normalize_location_key(location_name: str) -> str:
    return re.sub(r"\s+", " ", _clean_text(location_name).upper())


def _contains_location_pattern(normalized_location: str, pattern: str) -> bool:
    escaped = re.escape(pattern.upper())
    return re.search(rf"(^|[^A-Z]){escaped}([^A-Z]|$)", normalized_location) is not None


def _fallback_location_coordinates(location_name: str) -> dict[str, Any] | None:
    normalized_location = _normalize_location_key(location_name)
    for entry in LOCATION_COORDINATE_FALLBACKS:
        if any(_contains_location_pattern(normalized_location, pattern) for pattern in entry["patterns"]):
            return {
                "lat": float(entry["lat"]),
                "lon": float(entry["lon"]),
                "display_name": location_name,
            }
    return None


def _haversine_distance_km(start: dict[str, float], end: dict[str, float]) -> float:
    from math import asin, cos, radians, sin, sqrt

    radius_km = 6371
    d_lat = radians(end["lat"] - start["lat"])
    d_lon = radians(end["lon"] - start["lon"])
    lat1 = radians(start["lat"])
    lat2 = radians(end["lat"])
    a = sin(d_lat / 2) ** 2 + cos(lat1) * cos(lat2) * sin(d_lon / 2) ** 2
    return 2 * radius_km * asin(sqrt(a))


def _geocode_location(location_name: str) -> dict[str, Any] | None:
    queries = [
        location_name,
        f"{location_name}, India",
        f"{location_name}, Odisha, India",
    ]
    headers = {
        "User-Agent": "shipment-portal/1.0 (contact: local-app)",
        "Accept": "application/json",
    }
    for index, query in enumerate(queries):
        try:
            response = requests.get(
                NOMINATIM_SEARCH_URL,
                params={"q": query, "format": "jsonv2", "limit": 1},
                timeout=20,
                headers=headers,
            )
            if response.status_code != 200:
                continue
            payload = response.json()
            if isinstance(payload, list) and payload:
                first = payload[0]
                return {
                    "lat": float(first.get("lat")),
                    "lon": float(first.get("lon")),
                    "display_name": _clean_text(first.get("display_name")),
                }
        except requests.RequestException as exc:
            logger.warning("Location geocode request failed for %s: %s", location_name, exc)
            continue
        except (TypeError, ValueError) as exc:
            logger.warning("Location geocode parse failed for %s: %s", location_name, exc)
            continue
        finally:
            if index < len(queries) - 1:
                time.sleep(1.05)
    return None


def _route_distance_to_birgunj_km(lat: float, lon: float) -> float | None:
    try:
        response = requests.get(
            f"{OSRM_ROUTE_URL}/{lon},{lat};{BIRGUNJ_REFERENCE['lon']},{BIRGUNJ_REFERENCE['lat']}",
            params={"overview": "false"},
            timeout=20,
            headers={"User-Agent": "shipment-portal/1.0 (contact: local-app)"},
        )
        if response.status_code != 200:
            return None
        payload = response.json()
        routes = payload.get("routes") if isinstance(payload, dict) else None
        if not isinstance(routes, list) or not routes:
            return None
        distance_meters = routes[0].get("distance")
        if distance_meters is None:
            return None
        return round(float(distance_meters) / 1000, 2)
    except requests.RequestException as exc:
        logger.warning("Route distance lookup failed for (%s, %s): %s", lat, lon, exc)
        return None
    except (TypeError, ValueError) as exc:
        logger.warning("Route distance parse failed for (%s, %s): %s", lat, lon, exc)
        return None


def _resolve_location_distance(location_name: str) -> dict[str, Any]:
    location_text = _clean_text(location_name)
    if not location_text:
        return {"found": False, "distance_km": None, "display_name": ""}

    cache_key = _normalize_location_key(location_text)
    cache = _load_location_distance_cache()
    cached = cache.get(cache_key)
    fallback_coordinates = _fallback_location_coordinates(location_text)
    if isinstance(cached, dict) and (cached.get("found") or not fallback_coordinates):
        if fallback_coordinates and cached.get("distance_km") is not None:
            minimum_distance = round(
                _haversine_distance_km(
                    BIRGUNJ_REFERENCE,
                    {"lat": fallback_coordinates["lat"], "lon": fallback_coordinates["lon"]},
                ),
                2,
            )
            if float(cached.get("distance_km") or 0) >= minimum_distance:
                return cached
        elif not fallback_coordinates:
            return cached

    geocoded = _geocode_location(location_text) or fallback_coordinates
    if not geocoded:
        result = {"found": False, "distance_km": None, "display_name": location_text}
        cache[cache_key] = result
        _save_location_distance_cache(cache)
        return result

    haversine_distance = round(
        _haversine_distance_km(
            BIRGUNJ_REFERENCE,
            {"lat": geocoded["lat"], "lon": geocoded["lon"]},
        ),
        2,
    )
    route_distance = _route_distance_to_birgunj_km(geocoded["lat"], geocoded["lon"])
    if route_distance is None:
        route_distance = haversine_distance
    else:
        route_distance = max(round(route_distance, 2), haversine_distance)

    result = {
        "found": True,
        "distance_km": route_distance,
        "display_name": geocoded.get("display_name") or location_text,
    }
    cache[cache_key] = result
    _save_location_distance_cache(cache)
    return result


def _reconcile_customer_directory(db: Session) -> bool:
    entries = list(
        db.execute(
            select(CustomerDirectory).order_by(
                CustomerDirectory.source_count.desc(),
                CustomerDirectory.updated_at.desc(),
            )
        ).scalars()
    )
    touched = False
    for index, entry in enumerate(entries):
        if entry is None:
            continue
        for candidate in entries[index + 1 :]:
            if candidate is None:
                continue
            if not _is_customer_match(entry.display_name, candidate.display_name):
                continue
            winner = entry
            loser = candidate
            winner.display_name = _pick_better_display_name(winner.display_name, loser.display_name)
            merged_aliases = set(_aliases_from_json(winner.aliases_json))
            merged_aliases.update(_aliases_from_json(loser.aliases_json))
            merged_aliases.add(winner.display_name)
            merged_aliases.add(loser.display_name)
            winner.aliases_json = json.dumps(sorted(alias for alias in merged_aliases if alias), ensure_ascii=True)
            winner.source_count = (winner.source_count or 0) + (loser.source_count or 0)
            for shipment in db.execute(
                select(Shipment).where(Shipment.customer_name == loser.display_name)
            ).scalars():
                shipment.customer_name = winner.display_name
            db.delete(loser)
            entries[entries.index(candidate)] = None
            touched = True
    canonical_names = {
        entry.display_name
        for entry in db.execute(select(CustomerDirectory)).scalars()
    }
    for shipment in db.execute(select(Shipment)).scalars():
        canonical_customer = _sync_customer_directory(db, shipment.customer_name)
        if canonical_customer and shipment.customer_name != canonical_customer:
            shipment.customer_name = canonical_customer
            touched = True
        elif shipment.customer_name and shipment.customer_name not in canonical_names:
            shipment.customer_name = _format_customer_name(shipment.customer_name)
            touched = True
    return touched

def _shipment_to_dict(shipment: Shipment) -> dict[str, Any]:
    movement_category = _effective_shipment_movement(shipment)
    effective_location = _effective_shipment_location(shipment)
    action_required = _shipment_needs_action(shipment)
    source_type = _clean_text(getattr(shipment, "source_type", "")) or "manual"
    source_label = _clean_text(getattr(shipment, "source_label", "")) or ("Manual Entry" if source_type == "manual" else "")
    movement_diagnostics = _movement_diagnostics(
        shipment,
        movement_category,
        effective_location,
        action_required,
    )
    movement_since_date = _movement_since_date(
        movement_category,
        shipment.latest_time,
        shipment.port_arrival_date,
        getattr(shipment, "birgunj_arrival_date", ""),
        shipment.departure,
        getattr(shipment, "wagon_loaded_date", ""),
    )
    return {
        "id": shipment.id,
        "customer_name": _format_customer_name(shipment.customer_name),
        "container_number": shipment.container_number,
        "bl_number": shipment.bl_number,
        "shipment_status": shipment.shipment_status,
        "latest_location": effective_location,
        "latest_time": shipment.latest_time,
        "movement_since_date": movement_since_date,
        "port_arrival_date": shipment.port_arrival_date,
        "birgunj_arrival_date": getattr(shipment, "birgunj_arrival_date", ""),
        "pristine_booking_date": getattr(shipment, "pristine_booking_date", ""),
        "train_no": shipment.train_no,
        "departure": shipment.departure,
        "wagon_loaded_date": getattr(shipment, "wagon_loaded_date", ""),
        "concor_location_code": getattr(shipment, "concor_location_code", ""),
        "rail_status": shipment.rail_status,
        "movement_category": movement_category,
        "delay_days": shipment.delay_days,
        "wagon_no": shipment.wagon_no,
        "train_origin": shipment.train_origin,
        "train_destination": shipment.train_destination,
        "shipping_line": shipment.shipping_line,
        "tracking_source": shipment.tracking_source,
        "last_refresh_at": shipment.last_refresh_at,
        "last_refresh_status": shipment.last_refresh_status,
        "last_refresh_error": shipment.last_refresh_error,
        "clearance_doc_number": shipment.clearance_doc_number,
        "do_date": getattr(shipment, "do_date", ""),
        "document_status": getattr(shipment, "document_status", ""),
        "original_docs_received_date": getattr(shipment, "original_docs_received_date", ""),
        "action_required": action_required,
        "action_required_reason": "Pristine booking date is after Birgunj arrival" if action_required else "",
        "movement_diagnostics": movement_diagnostics,
        "source_type": source_type,
        "source_label": source_label,
        "source_batch_id": int(getattr(shipment, "source_batch_id", 0) or 0),
        "raw_source_row_id": int(getattr(shipment, "raw_source_row_id", 0) or 0),
    }


def _serialize_source_batches(
    db: Session,
    current_user: User,
    *,
    limit: int = 20,
) -> list[dict[str, Any]]:
    safe_limit = max(1, min(limit, 100))
    batches = list(
        db.execute(
            select(ShipmentBatch)
            .where(ShipmentBatch.user_id == current_user.id)
            .order_by(ShipmentBatch.created_at.desc(), ShipmentBatch.id.desc())
            .limit(safe_limit)
        ).scalars()
    )
    source_ids = sorted({batch.source_id for batch in batches if batch.source_id})
    source_map = (
        {
            source.id: source
            for source in db.execute(select(ShipmentSource).where(ShipmentSource.id.in_(source_ids))).scalars()
        }
        if source_ids
        else {}
    )
    return [
        {
            "id": batch.id,
            "source_id": batch.source_id,
            "batch_label": batch.batch_label,
            "source_sheet": batch.source_sheet,
            "header_row": batch.header_row,
            "status": batch.status,
            "imported_count": batch.imported_count,
            "duplicate_count": batch.duplicate_count,
            "invalid_count": batch.invalid_count,
            "skipped_blank_count": batch.skipped_blank_count,
            "created_at": batch.created_at.isoformat(),
            "source_label": source_map.get(batch.source_id).source_label if source_map.get(batch.source_id) else "",
            "source_type": source_map.get(batch.source_id).source_type if source_map.get(batch.source_id) else "",
            "source_reference": source_map.get(batch.source_id).source_reference if source_map.get(batch.source_id) else "",
        }
        for batch in batches
    ]


def _serialize_source_mappings(db: Session, current_user: User) -> list[dict[str, Any]]:
    profiles = list(
        db.execute(
            select(SourceMappingProfile)
            .where(SourceMappingProfile.user_id == current_user.id)
            .order_by(SourceMappingProfile.updated_at.desc(), SourceMappingProfile.id.desc())
        ).scalars()
    )
    return [
        {
            "id": profile.id,
            "source_type": profile.source_type,
            "source_sheet": profile.source_sheet,
            "profile_label": profile.profile_label,
            "mapping_json": _safe_json_loads(profile.mapping_json, {}),
            "updated_at": profile.updated_at.isoformat(),
        }
        for profile in profiles
    ]


def _serialize_source_connections(db: Session, current_user: User) -> list[dict[str, Any]]:
    connections = list(
        db.execute(
            select(SourceConnection)
            .where(SourceConnection.user_id == current_user.id)
            .order_by(SourceConnection.updated_at.desc(), SourceConnection.id.desc())
        ).scalars()
    )
    return [
        {
            "id": connection.id,
            "source_id": connection.source_id,
            "provider": connection.provider,
            "connection_label": connection.connection_label,
            "source_url": connection.source_url,
            "worksheet_name": connection.worksheet_name,
            "mapping_profile_id": connection.mapping_profile_id,
            "status": connection.status,
            "config": _safe_json_loads(connection.config_json, {}),
            "updated_at": connection.updated_at.isoformat(),
        }
        for connection in connections
    ]


def _ensure_shipment_columns() -> None:
    inspector = inspect(engine)
    existing_columns = {column["name"] for column in inspector.get_columns("shipments")}
    with engine.begin() as connection:
        for column_name, definition in SHIPMENT_COLUMN_DEFINITIONS.items():
            if column_name not in existing_columns:
                connection.execute(text(f"ALTER TABLE shipments ADD COLUMN {column_name} {definition}"))
        connection.execute(text("CREATE INDEX IF NOT EXISTS ix_shipments_status ON shipments (shipment_status)"))


def _json_dumps(value: Any) -> str:
    try:
        return json.dumps(value or {}, ensure_ascii=False)
    except (TypeError, ValueError):
        return "{}"


def _safe_json_loads(raw_value: Any, fallback: Any) -> Any:
    if raw_value in (None, ""):
        return fallback
    if isinstance(raw_value, (dict, list)):
        return raw_value
    try:
        return json.loads(raw_value)
    except (TypeError, json.JSONDecodeError):
        return fallback


def _get_or_create_source(
    db: Session,
    current_user: User,
    source_type: str,
    source_label: str,
    source_reference: str = "",
    metadata: dict[str, Any] | None = None,
) -> ShipmentSource:
    source = db.execute(
        select(ShipmentSource).where(
            ShipmentSource.user_id == current_user.id,
            ShipmentSource.source_type == source_type,
            ShipmentSource.source_label == source_label,
            ShipmentSource.source_reference == source_reference,
        )
    ).scalar_one_or_none()
    if source:
        if metadata:
            source.metadata_json = _json_dumps(metadata)
        source.status = "active"
        source.updated_at = datetime.utcnow()
        return source

    source = ShipmentSource(
        user_id=current_user.id,
        source_type=source_type,
        source_label=source_label,
        source_reference=source_reference,
        status="active",
        metadata_json=_json_dumps(metadata),
    )
    db.add(source)
    db.flush()
    return source


def _create_import_batch(
    db: Session,
    current_user: User,
    source: ShipmentSource,
    *,
    batch_label: str,
    source_sheet: str,
    header_row: int,
    mapping_json: dict[str, Any],
    imported_count: int,
    duplicate_count: int,
    invalid_count: int,
    skipped_blank_count: int,
) -> ShipmentBatch:
    batch = ShipmentBatch(
        user_id=current_user.id,
        source_id=source.id,
        batch_label=batch_label,
        source_sheet=source_sheet,
        header_row=header_row,
        mapping_json=_json_dumps(mapping_json),
        status="imported",
        imported_count=imported_count,
        duplicate_count=duplicate_count,
        invalid_count=invalid_count,
        skipped_blank_count=skipped_blank_count,
    )
    db.add(batch)
    db.flush()
    return batch


def _record_raw_import_rows(
    db: Session,
    current_user: User,
    source: ShipmentSource,
    batch: ShipmentBatch,
    normalized_rows: list[dict[str, Any]],
    customer_col: str,
    container_col: str,
    bl_col: str,
) -> dict[int, RawSourceRow]:
    rows_by_source_number: dict[int, RawSourceRow] = {}
    for row_obj in normalized_rows:
        source_row_number = int(row_obj.get("__source_row_number") or 0)
        container_number = _clean_container(row_obj.get(container_col))
        row_error = _container_format_error(container_number) if container_number else ""
        row_status = "blank" if not container_number else ("invalid" if row_error else "valid")
        mapped_row = {
            "customer_name": _clean_text(row_obj.get(customer_col)),
            "container_number": container_number,
            "bl_number": _normalize_bl_number(row_obj.get(bl_col)),
        }
        raw_row = {key: value for key, value in row_obj.items() if not str(key).startswith("__")}
        raw_source_row = RawSourceRow(
            user_id=current_user.id,
            source_id=source.id,
            batch_id=batch.id,
            source_row_number=source_row_number,
            row_status=row_status,
            row_error=row_error,
            raw_row_json=_json_dumps(raw_row),
            mapped_row_json=_json_dumps(mapped_row),
        )
        db.add(raw_source_row)
        db.flush()
        rows_by_source_number[source_row_number] = raw_source_row
    return rows_by_source_number


def _load_recent_mapping_for_sheet(
    db: Session,
    current_user: User,
    sheet_name: str,
    source_type: str = "excel_upload",
) -> dict[str, Any]:
    normalized_sheet = _clean_text(sheet_name)
    if normalized_sheet:
        profile = db.execute(
            select(SourceMappingProfile)
            .where(
                SourceMappingProfile.user_id == current_user.id,
                SourceMappingProfile.source_type == source_type,
                SourceMappingProfile.source_sheet == normalized_sheet,
            )
            .order_by(SourceMappingProfile.updated_at.desc(), SourceMappingProfile.id.desc())
            .limit(1)
        ).scalar_one_or_none()
        if profile:
            try:
                return json.loads(profile.mapping_json or "{}")
            except (TypeError, json.JSONDecodeError):
                return {}

    recent_profile = db.execute(
        select(SourceMappingProfile)
        .where(
            SourceMappingProfile.user_id == current_user.id,
            SourceMappingProfile.source_type == source_type,
        )
        .order_by(SourceMappingProfile.updated_at.desc(), SourceMappingProfile.id.desc())
        .limit(1)
    ).scalar_one_or_none()
    if not recent_profile:
        return {}
    try:
        return json.loads(recent_profile.mapping_json or "{}")
    except (TypeError, json.JSONDecodeError):
        return {}


def _upsert_mapping_profile(
    db: Session,
    current_user: User,
    *,
    source_type: str,
    source_sheet: str,
    profile_label: str,
    mapping_json: dict[str, Any],
) -> SourceMappingProfile:
    normalized_sheet = _clean_text(source_sheet)
    profile = db.execute(
        select(SourceMappingProfile).where(
            SourceMappingProfile.user_id == current_user.id,
            SourceMappingProfile.source_type == source_type,
            SourceMappingProfile.source_sheet == normalized_sheet,
        )
    ).scalar_one_or_none()
    if profile:
        profile.profile_label = profile_label or profile.profile_label
        profile.mapping_json = _json_dumps(mapping_json)
        profile.updated_at = datetime.utcnow()
        return profile

    profile = SourceMappingProfile(
        user_id=current_user.id,
        source_type=source_type,
        source_sheet=normalized_sheet,
        profile_label=profile_label or normalized_sheet or source_type,
        mapping_json=_json_dumps(mapping_json),
    )
    db.add(profile)
    db.flush()
    return profile


def _parse_google_sheet_reference(source_url: str) -> dict[str, str]:
    normalized_url = _clean_text(source_url)
    if not normalized_url:
        raise HTTPException(status_code=400, detail="Google Sheets URL is required")

    patterns = [
        r"/spreadsheets/d/([a-zA-Z0-9-_]+)",
        r"[?&]id=([a-zA-Z0-9-_]+)",
    ]
    sheet_id = ""
    for pattern in patterns:
        match = re.search(pattern, normalized_url)
        if match:
            sheet_id = match.group(1)
            break
    if not sheet_id:
        raise HTTPException(status_code=400, detail="Could not detect a Google Sheet ID from the provided URL")

    gid_match = re.search(r"[?&#]gid=(\d+)", normalized_url)
    gid = gid_match.group(1) if gid_match else ""
    return {"sheet_id": sheet_id, "gid": gid}


def _download_google_sheet_workbook(source_url: str) -> tuple[bytes, dict[str, str]]:
    parsed = _parse_google_sheet_reference(source_url)
    export_url = f"https://docs.google.com/spreadsheets/d/{parsed['sheet_id']}/export?format=xlsx"
    try:
        response = requests.get(
            export_url,
            timeout=(4, 20),
            allow_redirects=True,
            headers={"User-Agent": "Mozilla/5.0", "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*"},
        )
    except requests.RequestException as exc:
        raise HTTPException(status_code=400, detail=f"Unable to fetch Google Sheet: {exc}") from exc

    if response.status_code != 200 or not response.content:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unable to fetch Google Sheet. The current read-only connection works only for sheets that are "
                "accessible by link. For private sheets, the next secure step is Google sign-in with read-only access."
            ),
        )
    content_type = _clean_text(response.headers.get("content-type", "")).lower()
    if "spreadsheetml" not in content_type and not response.content.startswith(b"PK"):
        raise HTTPException(
            status_code=400,
            detail=(
                "Google Sheet could not be downloaded as an Excel workbook. The current connection expects a "
                "link-accessible sheet; private-sheet sync will require Google sign-in."
            ),
        )
    return response.content, parsed


def _sheet_title_from_workbook(workbook, parsed_reference: dict[str, str]) -> str:
    title = _clean_text(getattr(getattr(workbook, "properties", None), "title", ""))
    if title:
        return title
    return f"Google Sheet {parsed_reference.get('sheet_id', '')[:8]}".strip()


def _fetch_google_sheet_title(source_url: str) -> str:
    try:
        response = requests.get(
            source_url,
            timeout=(4, 12),
            allow_redirects=True,
            headers={"User-Agent": "Mozilla/5.0", "Accept": "text/html,application/xhtml+xml"},
        )
    except requests.RequestException:
        return ""
    if response.status_code != 200 or not response.text:
        return ""
    text = response.text
    og_match = re.search(r'<meta[^>]+property="og:title"[^>]+content="([^"]+)"', text, re.IGNORECASE)
    if og_match:
        return _clean_text(unescape(og_match.group(1)))
    title_match = re.search(r"<title>(.*?)</title>", text, re.IGNORECASE | re.DOTALL)
    if not title_match:
        return ""
    title = _clean_text(unescape(title_match.group(1)))
    if title.lower().endswith("- google sheets"):
        title = _clean_text(title[: -len("- google sheets")])
    return title


def _resolve_default_user_id(db: Session) -> int:
    user = db.execute(select(User).where(User.email == DEFAULT_LOCAL_USER_EMAIL)).scalar_one_or_none()
    if user:
        return user.id
    user = db.execute(select(User).order_by(User.id.asc())).scalars().first()
    if user:
        return user.id
    fallback_user = User(email=DEFAULT_LOCAL_USER_EMAIL, password_hash=hash_password(settings.demo_password))
    db.add(fallback_user)
    db.commit()
    db.refresh(fallback_user)
    return fallback_user.id


def _apply_legacy_row(db: Session, shipment: Shipment, row: dict[str, Any], default_user_id: int) -> bool:
    changed = False
    canonical_customer = _sync_customer_directory(db, _clean_text(row.get("customer_name")))
    field_updates = {
        "user_id": shipment.user_id or default_user_id,
        "customer_name": canonical_customer,
        "container_number": _clean_container(row.get("container_number")),
        "bl_number": _clean_bl(row.get("bl_number")),
        "shipment_status": _clean_text(row.get("shipment_status")).lower() or "active",
        "latest_location": _clean_text(row.get("latest_location")),
        "latest_time": _clean_text(row.get("latest_time")),
        "train_no": _clean_text(row.get("train_no")),
        "departure": _clean_text(row.get("departure")),
        "rail_status": _clean_text(row.get("rail_status")),
        "movement_category": _normalize_existing_movement(_clean_text(row.get("movement_category")) or "Hi Seas"),
        "delay_days": float(row.get("delay_days") or 0),
        "wagon_no": _clean_text(row.get("wagon_no")),
        "train_origin": _clean_text(row.get("train_origin")),
        "train_destination": _clean_text(row.get("train_destination")),
        "shipping_line": _clean_text(row.get("shipping_line")),
        "tracking_source": _clean_text(row.get("tracking_source")),
        "last_refresh_at": _clean_text(row.get("last_refresh_at")),
        "last_refresh_status": _clean_text(row.get("last_refresh_status")),
        "last_refresh_error": _clean_text(row.get("last_refresh_error")),
        "clearance_doc_number": _clean_text(row.get("clearance_doc_number")),
    }
    for field_name, next_value in field_updates.items():
        if getattr(shipment, field_name) != next_value:
            setattr(shipment, field_name, next_value)
            changed = True
    return changed


def _bootstrap_shipments_from_legacy_state(db: Session) -> None:
    legacy_rows = _load_state().get("shipments", [])
    default_user_id = _resolve_default_user_id(db)
    existing_shipments = db.execute(select(Shipment)).scalars().all()
    shipments_by_key = {
        (shipment.container_number, _normalize_bl_number(shipment.bl_number), _format_customer_name(shipment.customer_name)): shipment
        for shipment in existing_shipments
    }
    existing_ids = {shipment.id for shipment in existing_shipments}
    touched = False
    for row in legacy_rows:
        container_number = _clean_container(row.get("container_number"))
        if not container_number:
            continue
        canonical_customer = _sync_customer_directory(db, _clean_text(row.get("customer_name")))
        row_key = (container_number, _normalize_bl_number(row.get("bl_number")), canonical_customer)
        shipment = shipments_by_key.get(row_key)
        if shipment is None:
            shipment = Shipment(user_id=default_user_id, customer_name="", container_number=container_number, bl_number="")
            row_id = row.get("id")
            if isinstance(row_id, int) and row_id not in existing_ids:
                shipment.id = row_id
                existing_ids.add(row_id)
            db.add(shipment)
            shipments_by_key[row_key] = shipment
            touched = True
        if _apply_legacy_row(db, shipment, row, default_user_id):
            touched = True
    if _reconcile_customer_directory(db):
        touched = True
    if touched:
        db.commit()


def _ensure_storage_ready(db: Session) -> None:
    global _storage_ready
    if _storage_ready:
        return
    _ensure_shipment_columns()
    _bootstrap_shipments_from_legacy_state(db)
    _storage_ready = True


def _adopt_demo_shipments_for_user(db: Session, current_user: User) -> None:
    if not settings.enable_demo_shipment_adoption:
        return
    if current_user.email == DEFAULT_LOCAL_USER_EMAIL:
        return
    user_has_shipments = db.execute(
        select(func.count()).select_from(Shipment).where(Shipment.user_id == current_user.id)
    ).scalar_one()
    if user_has_shipments:
        return
    demo_user = db.execute(select(User).where(User.email == DEFAULT_LOCAL_USER_EMAIL)).scalar_one_or_none()
    if not demo_user or demo_user.id == current_user.id:
        return
    other_owner_count = db.execute(
        select(func.count())
        .select_from(Shipment)
        .where(Shipment.user_id.not_in([demo_user.id, current_user.id]))
    ).scalar_one()
    if other_owner_count:
        return
    demo_shipments = list(db.execute(select(Shipment).where(Shipment.user_id == demo_user.id)).scalars())
    if not demo_shipments:
        return
    for shipment in demo_shipments:
        shipment.user_id = current_user.id
    db.commit()


def _ensure_user_scope_ready(db: Session, current_user: User) -> None:
    _ensure_storage_ready(db)
    _adopt_demo_shipments_for_user(db, current_user)
    _reconcile_bl_number_normalization_for_user(db, current_user)
    _reconcile_orphan_shipments_for_user(db, current_user)


def _user_shipment_select(current_user: User):
    return select(Shipment).where(Shipment.user_id == current_user.id)


def _user_from_access_token(db: Session, access_token: str) -> User:
    token = _clean_text(access_token)
    if not token:
        raise HTTPException(status_code=401, detail="Missing authentication token")
    try:
        payload = jwt.decode(token, settings.secret_key, algorithms=["HS256"])
        user_id = int(payload.get("sub"))
    except (JWTError, ValueError, TypeError):
        raise HTTPException(status_code=401, detail="Invalid authentication token")
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    return user


def get_portal_user(
    db: Session = Depends(get_db),
    authorization: str | None = Header(default=None),
) -> User:
    auth_text = _clean_text(authorization)
    if auth_text.lower().startswith("bearer "):
        return _user_from_access_token(db, auth_text.split(" ", 1)[1])
    if not settings.allow_demo_portal_fallback:
        raise HTTPException(status_code=401, detail="Missing authentication token")
    default_user_id = _resolve_default_user_id(db)
    fallback_user = db.get(User, default_user_id)
    if not fallback_user:
        raise HTTPException(status_code=401, detail="User not found")
    return fallback_user


def _build_tracking_payload_from_sources(
    container_number: str,
    shipment: Shipment | None,
    source_payload: dict[str, Any],
) -> dict[str, Any]:
    ldb_data = source_payload.get("ldb") or {}
    concor_data = source_payload.get("concor") or {}
    pristine_data = source_payload.get("pristine") or {}
    resolved_state = resolve_shipment_state(
        shipment,
        ldb_data,
        concor_data,
        pristine_data,
        now=datetime.now(APP_TIMEZONE).replace(tzinfo=None),
    )

    delay_reference = _movement_since_date(
        resolved_state.movement_category,
        resolved_state.latest_time,
        resolved_state.port_arrival_date,
        resolved_state.birgunj_arrival_date,
        resolved_state.departure,
        resolved_state.wagon_loaded_date,
    )
    delay_days = _compute_delay_days(delay_reference)

    data = {
        "latest_location": resolved_state.latest_location,
        "latest_time": resolved_state.latest_time,
        "port_arrival_date": resolved_state.port_arrival_date,
        "birgunj_arrival_date": resolved_state.birgunj_arrival_date,
        "pristine_booking_date": resolved_state.pristine_booking_date,
        "train_no": resolved_state.train_no,
        "departure": resolved_state.departure,
        "wagon_loaded_date": resolved_state.wagon_loaded_date,
        "concor_location_code": resolved_state.concor_location_code,
        "rail_status": resolved_state.rail_status,
        "movement_category": resolved_state.movement_category,
        "delay_days": delay_days,
        "wagon_no": concor_data.get("wagon_no", ""),
        "train_origin": concor_data.get("train_origin", ""),
        "train_destination": concor_data.get("train_destination", ""),
        "shipping_line": concor_data.get("shipping_line", ""),
        "tracking_source": resolved_state.tracking_source,
        "movement_diagnostics": resolved_state.decision_trace,
        "completion_frozen": resolved_state.completion_frozen,
    }

    return {
        "data": data,
        "status": source_payload.get("status", "success"),
        "error": source_payload.get("error", "") or "",
        "cached": bool(source_payload.get("cached")),
        "has_data": bool(source_payload.get("has_data")),
    }


def _build_tracking_payload(
    container_number: str,
    shipment: Shipment | None = None,
    use_cache: bool = True,
) -> dict[str, Any]:
    container_number = _clean_container(container_number)
    if not container_number:
        return {
            "data": {},
            "status": "error",
            "error": "Missing container number",
            "cached": False,
            "has_data": False,
        }
    source_payload = _fetch_tracking_sources(container_number, use_cache=use_cache)
    return _build_tracking_payload_from_sources(container_number, shipment, source_payload)


def _apply_tracking_payload(shipment: Shipment, payload: dict[str, Any]) -> Shipment:
    data = payload.get("data") or {}
    shipment.latest_location = data.get("latest_location", "") or ""
    shipment.latest_time = data.get("latest_time", "") or ""
    shipment.port_arrival_date = data.get("port_arrival_date", "") or ""
    shipment.birgunj_arrival_date = data.get("birgunj_arrival_date", "") or ""
    shipment.pristine_booking_date = data.get("pristine_booking_date", "") or ""
    shipment.train_no = data.get("train_no", "") or ""
    shipment.departure = data.get("departure", "") or ""
    shipment.wagon_loaded_date = data.get("wagon_loaded_date", "") or ""
    shipment.concor_location_code = data.get("concor_location_code", "") or ""
    shipment.rail_status = data.get("rail_status", "") or ""
    shipment.movement_category = _normalize_existing_movement(data.get("movement_category", "") or "Hi Seas")
    shipment.delay_days = float(data.get("delay_days", 0) or 0)
    shipment.wagon_no = data.get("wagon_no", "") or ""
    shipment.train_origin = data.get("train_origin", "") or ""
    shipment.train_destination = data.get("train_destination", "") or ""
    shipment.shipping_line = data.get("shipping_line", "") or ""
    shipment.tracking_source = data.get("tracking_source", "") or ""
    shipment.last_refresh_at = _now_datetime()
    shipment.last_refresh_status = payload.get("status", "success")
    shipment.last_refresh_error = payload.get("error", "") or ""
    if data.get("movement_diagnostics"):
        setattr(shipment, "_movement_diagnostics", data.get("movement_diagnostics"))
    return shipment


def _refresh_one_shipment(shipment: Shipment, use_cache: bool = True) -> Shipment:
    container_number = _clean_container(shipment.container_number)
    if not container_number:
        shipment.last_refresh_status = "error"
        shipment.last_refresh_error = "Missing container number"
        shipment.last_refresh_at = _now_datetime()
        return shipment
    return _apply_tracking_payload(shipment, _build_tracking_payload(container_number, shipment=shipment, use_cache=use_cache))


def _get_shipments(db: Session, current_user: User) -> list[Shipment]:
    return list(
        db.execute(_user_shipment_select(current_user).order_by(Shipment.updated_at.desc(), Shipment.id.desc())).scalars()
    )


def _refresh_active_shipments_for_user(
    db: Session,
    user_id: int,
    *,
    task_id: str | None = None,
    audit_action: str | None = None,
) -> dict[str, int]:
    shipments = list(
        db.execute(
            select(Shipment).where(
                Shipment.user_id == user_id,
                Shipment.shipment_status == "active",
            )
        ).scalars()
    )
    refreshed_group_count = len(_group_dashboard_rows(shipments))
    unique_containers = sorted(
        {
            _clean_container(shipment.container_number)
            for shipment in shipments
            if _clean_container(shipment.container_number)
        }
    )
    total_containers = len(unique_containers)

    if task_id:
        _set_refresh_job(
            task_id,
            state="running",
            total=total_containers,
            completed=0,
            progress=0,
            message=(
                f"Refreshing live tracking for {refreshed_group_count} shipment"
                f"{'s' if refreshed_group_count != 1 else ''}."
            ),
        )

    if not unique_containers:
        if audit_action:
            _log_audit_event(
                db,
                user_id,
                audit_action,
                shipment_status="active",
                details={
                    "refreshed_count": refreshed_group_count,
                    "refreshed_container_count": 0,
                },
            )
        db.commit()
        if task_id:
            _set_refresh_job(
                task_id,
                state="completed",
                completed=0,
                progress=100,
                message="No active shipments needed refreshing.",
                refreshed_count=refreshed_group_count,
                refreshed_container_count=0,
            )
        return {
            "refreshed_count": refreshed_group_count,
            "refreshed_container_count": 0,
        }

    resolved_payloads: dict[str, dict[str, Any]] = {}
    with ThreadPoolExecutor(max_workers=max(1, min(TRACKING_POOL_WORKERS, total_containers))) as executor:
        futures = {
            executor.submit(lambda value: _build_tracking_payload(value, use_cache=False), container_number): container_number
            for container_number in unique_containers
        }
        completed = 0
        for future in as_completed(futures):
            container_number = futures[future]
            try:
                resolved_payloads[container_number] = future.result()
            except Exception:
                logger.exception(
                    "Tracking payload build failed during bulk refresh for container %s",
                    container_number,
                )
                resolved_payloads[container_number] = {
                    "ldb": {},
                    "concor": {},
                    "pristine": {},
                    "status": "error",
                    "error": "Tracking fetch failed",
                    "cached": False,
                    "has_data": False,
                }
            completed += 1
            if task_id:
                _set_refresh_job(
                    task_id,
                    completed=completed,
                    progress=round((completed / total_containers) * 100),
                    message="Refreshing live tracking.",
                )

    for shipment in shipments:
        container_number = _clean_container(shipment.container_number)
        if container_number in resolved_payloads:
            _apply_tracking_payload(shipment, resolved_payloads[container_number])

    if audit_action:
        _log_audit_event(
            db,
            user_id,
            audit_action,
            shipment_status="active",
            details={
                "refreshed_count": refreshed_group_count,
                "refreshed_container_count": total_containers,
            },
        )

    db.commit()
    if task_id:
        _set_refresh_job(
            task_id,
            state="completed",
            completed=total_containers,
            progress=100,
            message=(
                f"Live tracking updated for {refreshed_group_count} shipment"
                f"{'s' if refreshed_group_count != 1 else ''}."
            ),
            refreshed_count=refreshed_group_count,
            refreshed_container_count=total_containers,
        )
    return {
        "refreshed_count": refreshed_group_count,
        "refreshed_container_count": total_containers,
    }


def _set_refresh_job(task_id: str, **updates: Any) -> dict[str, Any]:
    with _refresh_jobs_lock:
        jobs = _cleanup_refresh_jobs(_load_refresh_jobs())
        _refresh_jobs.clear()
        _refresh_jobs.update(jobs)
        job = _refresh_jobs.setdefault(task_id, {})
        job.update(updates)
        job["updated_at"] = datetime.now(APP_TIMEZONE).isoformat()
        _refresh_jobs[task_id] = job
        _save_refresh_jobs(_refresh_jobs)
        return dict(job)


def _get_refresh_job(task_id: str) -> dict[str, Any] | None:
    with _refresh_jobs_lock:
        jobs = _cleanup_refresh_jobs(_load_refresh_jobs())
        _refresh_jobs.clear()
        _refresh_jobs.update(jobs)
        job = _normalize_refresh_job(_refresh_jobs.get(task_id))
        if job:
            _refresh_jobs[task_id] = job
            _save_refresh_jobs(_refresh_jobs)
            return dict(job)
        return None


def _run_background_refresh_cycle() -> None:
    connect_args = {"check_same_thread": False} if normalized_database_url.startswith("sqlite") else {}
    job_engine = create_engine(normalized_database_url, future=True, connect_args=connect_args)
    from sqlalchemy.orm import sessionmaker

    JobSessionLocal = sessionmaker(bind=job_engine, autoflush=False, autocommit=False, future=True)
    db = JobSessionLocal()
    try:
        user_ids = [
            row[0]
            for row in db.execute(
                select(Shipment.user_id)
                .where(Shipment.shipment_status == "active")
                .distinct()
            ).all()
            if row and row[0]
        ]
        refreshed_users = 0
        refreshed_shipments = 0
        for user_id in user_ids:
            _update_background_refresh_lease(state="running", current_user_id=user_id)
            try:
                result = _refresh_active_shipments_for_user(db, user_id, audit_action=None)
            except Exception:
                logger.exception(
                    "Background refresh failed for user %s during scheduled cycle",
                    user_id,
                )
                db.rollback()
                continue
            if result["refreshed_count"] > 0:
                refreshed_users += 1
                refreshed_shipments += result["refreshed_count"]
        _update_background_refresh_lease(
            state="idle",
            current_user_id=0,
            refreshed_users=refreshed_users,
            refreshed_shipments=refreshed_shipments,
        )
    finally:
        db.close()
        job_engine.dispose()


def _run_refresh_all_job(task_id: str, db_url: str, user_id: int) -> None:
    from sqlalchemy.orm import sessionmaker

    connect_args = {"check_same_thread": False} if db_url.startswith("sqlite") else {}
    job_engine = create_engine(db_url, future=True, connect_args=connect_args)
    JobSessionLocal = sessionmaker(bind=job_engine, autoflush=False, autocommit=False, future=True)
    db = JobSessionLocal()
    try:
        _refresh_active_shipments_for_user(
            db,
            user_id,
            task_id=task_id,
            audit_action="shipment_all_refreshed",
        )
    except Exception as exc:
        logger.exception("Refresh-all job %s failed for user %s", task_id, user_id)
        db.rollback()
        _set_refresh_job(
            task_id,
            state="failed",
            progress=100,
            message=str(exc) or "Refresh failed.",
        )
    finally:
        db.close()
        job_engine.dispose()


def _resolve_group_shipments(
    db: Session,
    current_user: User,
    bl_number: Any = None,
    container_numbers: list[Any] | None = None,
    include_archived: bool = False,
) -> list[Shipment]:
    normalized_bl = _normalize_bl_number(bl_number)
    base_query = _user_shipment_select(current_user)
    if not include_archived:
        base_query = base_query.where(Shipment.shipment_status != "archived")

    if normalized_bl:
        return [
            shipment
            for shipment in db.execute(base_query).scalars()
            if _normalize_bl_number(shipment.bl_number) == normalized_bl
        ]

    cleaned_containers = [
        _clean_container(value)
        for value in (container_numbers or [])
        if _clean_container(value)
    ]
    if not cleaned_containers:
        return []

    matched_shipments = list(
        db.execute(base_query.where(Shipment.container_number.in_(cleaned_containers))).scalars()
    )
    derived_bl = _first_non_empty([shipment.bl_number for shipment in matched_shipments])
    normalized_derived_bl = _normalize_bl_number(derived_bl)
    if normalized_derived_bl:
        return [
            shipment
            for shipment in db.execute(base_query).scalars()
            if _normalize_bl_number(shipment.bl_number) == normalized_derived_bl
        ]
    return matched_shipments


def _apply_group_status_transition(
    db: Session,
    current_user: User,
    *,
    bl_number: Any = None,
    container_numbers: list[Any] | None = None,
    next_status: str,
    clearance_doc_number: str = "",
) -> tuple[list[Shipment], str]:
    shipments = _resolve_group_shipments(
        db,
        current_user,
        bl_number=bl_number,
        container_numbers=container_numbers or [],
        include_archived=next_status == "active",
    )
    if not shipments:
        raise HTTPException(status_code=404, detail="Shipment group not found")

    if next_status == "active":
        if any((shipment.shipment_status or "").lower() == "completed" for shipment in shipments):
            for shipment in shipments:
                if (shipment.shipment_status or "").lower() == "completed":
                    shipment.shipment_status = "active"
            return shipments, _first_non_empty([shipment.clearance_doc_number for shipment in shipments]) or clearance_doc_number

        target_bl = _normalize_bl_number(_first_non_empty([shipment.bl_number for shipment in shipments]))
        restore_containers = sorted(
            {
                _clean_container(shipment.container_number)
                for shipment in shipments
                if _clean_container(shipment.container_number)
            }
        )
        shipments_by_id: dict[int, Shipment] = {}
        archived_candidates = list(
            db.execute(
                _user_shipment_select(current_user).where(
                    Shipment.shipment_status == "archived",
                )
            ).scalars()
        )
        for candidate in archived_candidates:
            candidate_container = _clean_container(candidate.container_number)
            candidate_bl = _normalize_bl_number(candidate.bl_number)
            if target_bl:
                if candidate_bl == target_bl:
                    shipments_by_id[candidate.id] = candidate
                elif not candidate_bl and candidate_container in restore_containers:
                    shipments_by_id[candidate.id] = candidate
            elif not candidate_bl and candidate_container in restore_containers:
                shipments_by_id[candidate.id] = candidate
        shipments = list(shipments_by_id.values())
        if not shipments:
            raise HTTPException(status_code=404, detail="Archived shipment group not found")

    if next_status == "archived":
        existing_doc_number = _first_non_empty([shipment.clearance_doc_number for shipment in shipments])
        if not existing_doc_number:
            raise HTTPException(
                status_code=400,
                detail="Complete the BL with a clearance document number before archiving.",
            )
        target_bl = _normalize_bl_number(_first_non_empty([shipment.bl_number for shipment in shipments]))
        archive_containers = sorted(
            {
                _clean_container(shipment.container_number)
                for shipment in shipments
                if _clean_container(shipment.container_number)
            }
        )
        related_active_shipments = list(
            db.execute(
                _user_shipment_select(current_user).where(
                    Shipment.shipment_status != "archived",
                    Shipment.container_number.in_(archive_containers),
                )
            ).scalars()
        )
        shipments_by_id = {shipment.id: shipment for shipment in shipments}
        for related in related_active_shipments:
            related_bl = _normalize_bl_number(related.bl_number)
            if related_bl == target_bl or not related_bl:
                shipments_by_id.setdefault(related.id, related)
        shipments = list(shipments_by_id.values())

    effective_doc_number = _first_non_empty([shipment.clearance_doc_number for shipment in shipments]) or clearance_doc_number
    if next_status == "completed" and not effective_doc_number:
        raise HTTPException(status_code=400, detail="clearance_doc_number is required to complete a shipment")

    for shipment in shipments:
        shipment.shipment_status = next_status
        if next_status == "completed":
            shipment.clearance_doc_number = effective_doc_number

    _log_audit_event(
        db,
        current_user.id,
        "shipment_group_restored" if next_status == "active" else "shipment_group_status_updated",
        bl_number=_first_non_empty([shipment.bl_number for shipment in shipments]),
        container_number=_first_non_empty([shipment.container_number for shipment in shipments]),
        shipment_status=next_status,
        details={
            "count": len(shipments),
            "container_numbers": [shipment.container_number for shipment in shipments],
            "clearance_doc_number": effective_doc_number,
        },
    )
    return shipments, effective_doc_number


def _delete_group_shipments_internal(
    db: Session,
    current_user: User,
    *,
    bl_number: Any = None,
    container_numbers: list[Any] | None = None,
) -> list[Shipment]:
    shipments = _resolve_group_shipments(
        db,
        current_user,
        bl_number=bl_number,
        container_numbers=container_numbers or [],
        include_archived=True,
    )
    if not shipments:
        raise HTTPException(status_code=404, detail="Shipment group not found")
    _log_audit_event(
        db,
        current_user.id,
        "shipment_group_deleted",
        bl_number=_first_non_empty([shipment.bl_number for shipment in shipments]),
        container_number=_first_non_empty([shipment.container_number for shipment in shipments]),
        shipment_status="deleted",
        details={
            "count": len(shipments),
            "container_numbers": [shipment.container_number for shipment in shipments],
        },
    )
    for shipment in shipments:
        db.delete(shipment)
    return shipments


def _first_non_empty(values: list[str]) -> str:
    for value in values:
        text_value = _clean_text(value)
        if text_value:
            return text_value
    return ""


def _shipment_identity_key(shipment: Shipment | dict[str, Any]) -> tuple[str, str]:
    if isinstance(shipment, Shipment):
        customer_name = shipment.customer_name
        container_number = shipment.container_number
    else:
        customer_name = shipment.get("customer_name", "")
        container_number = shipment.get("container_number", "")
    return _format_customer_name(customer_name), _clean_container(container_number)


def _reconcile_orphan_shipments_for_user(db: Session, current_user: User) -> None:
    shipments = list(db.execute(_user_shipment_select(current_user)).scalars())
    if not shipments:
        return

    valid_by_key: dict[tuple[str, str], list[Shipment]] = {}
    orphan_rows: list[Shipment] = []
    for shipment in shipments:
        identity_key = _shipment_identity_key(shipment)
        if not identity_key[1]:
            continue
        if _normalize_bl_number(shipment.bl_number):
            valid_by_key.setdefault(identity_key, []).append(shipment)
        else:
            orphan_rows.append(shipment)

    deleted_count = 0
    for orphan in orphan_rows:
        matches = valid_by_key.get(_shipment_identity_key(orphan), [])
        if len(matches) == 1:
            db.delete(orphan)
            deleted_count += 1

    if deleted_count:
        _log_audit_event(
            db,
            current_user.id,
            "shipment_orphans_reconciled",
            shipment_status="active",
            details={"deleted_count": deleted_count},
        )
        db.commit()


def _shipment_record_identity(shipment: Shipment) -> tuple[str, str, str, str]:
    return (
        _clean_container(shipment.container_number),
        _normalize_bl_number(shipment.bl_number),
        _format_customer_name(shipment.customer_name),
        _clean_text(shipment.shipment_status).lower(),
    )


def _shipment_quality_score(shipment: Shipment) -> int:
    fields = [
        shipment.latest_location,
        shipment.latest_time,
        shipment.port_arrival_date,
        shipment.birgunj_arrival_date,
        shipment.pristine_booking_date,
        shipment.train_no,
        shipment.departure,
        shipment.rail_status,
        shipment.movement_category,
        shipment.wagon_no,
        shipment.train_origin,
        shipment.train_destination,
        shipment.shipping_line,
        shipment.tracking_source,
        shipment.last_refresh_at,
        shipment.last_refresh_status,
        shipment.last_refresh_error,
        shipment.clearance_doc_number,
        shipment.source_type,
        shipment.source_label,
    ]
    score = sum(1 for value in fields if _clean_text(value))
    if shipment.source_batch_id:
        score += 1
    if shipment.raw_source_row_id:
        score += 1
    return score


def _merge_shipment_record(preferred: Shipment, duplicate: Shipment) -> None:
    merge_fields = [
        "latest_location",
        "latest_time",
        "port_arrival_date",
        "birgunj_arrival_date",
        "pristine_booking_date",
        "train_no",
        "departure",
        "rail_status",
        "movement_category",
        "wagon_no",
        "train_origin",
        "train_destination",
        "shipping_line",
        "tracking_source",
        "last_refresh_at",
        "last_refresh_status",
        "last_refresh_error",
        "clearance_doc_number",
        "source_type",
        "source_label",
    ]
    for field_name in merge_fields:
        if not _clean_text(getattr(preferred, field_name, "")) and _clean_text(getattr(duplicate, field_name, "")):
            setattr(preferred, field_name, getattr(duplicate, field_name))
    if not preferred.source_batch_id and duplicate.source_batch_id:
        preferred.source_batch_id = duplicate.source_batch_id
    if not preferred.raw_source_row_id and duplicate.raw_source_row_id:
        preferred.raw_source_row_id = duplicate.raw_source_row_id


def _reconcile_bl_number_normalization_for_user(db: Session, current_user: User) -> None:
    shipments = list(db.execute(_user_shipment_select(current_user)).scalars())
    if not shipments:
        return

    updated_count = 0
    for shipment in shipments:
        normalized_bl = _normalize_bl_number(shipment.bl_number)
        if normalized_bl != _clean_text(shipment.bl_number):
            shipment.bl_number = normalized_bl
            updated_count += 1

    shipments = sorted(shipments, key=lambda shipment: (_shipment_record_identity(shipment), -_shipment_quality_score(shipment), shipment.id))
    deleted_count = 0
    kept_by_identity: dict[tuple[str, str, str, str], Shipment] = {}

    for shipment in shipments:
        identity = _shipment_record_identity(shipment)
        keeper = kept_by_identity.get(identity)
        if not identity[0]:
            continue
        if keeper is None:
            kept_by_identity[identity] = shipment
            continue
        preferred = keeper
        duplicate = shipment
        if _shipment_quality_score(duplicate) > _shipment_quality_score(preferred):
            preferred, duplicate = duplicate, preferred
            kept_by_identity[identity] = preferred
        _merge_shipment_record(preferred, duplicate)
        db.delete(duplicate)
        deleted_count += 1

    if updated_count or deleted_count:
        _log_audit_event(
            db,
            current_user.id,
            "shipment_orphans_reconciled",
            shipment_status="active",
            details={"normalized_bl_count": updated_count, "deleted_count": deleted_count},
        )
        db.commit()


def _group_dashboard_rows(shipments: list[Shipment]) -> list[dict[str, Any]]:
    suppressed_orphan_keys: set[tuple[str, str]] = set()
    archived_group_keys = {
        (_clean_container(shipment.container_number), _normalize_bl_number(shipment.bl_number))
        for shipment in shipments
        if shipment.shipment_status == "archived"
        and _clean_container(shipment.container_number)
        and _normalize_bl_number(shipment.bl_number)
    }
    archived_blank_containers = {
        _clean_container(shipment.container_number)
        for shipment in shipments
        if shipment.shipment_status == "archived"
        and _clean_container(shipment.container_number)
        and not _normalize_bl_number(shipment.bl_number)
    }
    for shipment in shipments:
        if shipment.shipment_status == "archived":
            continue
        if _normalize_bl_number(shipment.bl_number):
            suppressed_orphan_keys.add(_shipment_identity_key(shipment))

    grouped: dict[str, list[dict[str, Any]]] = {}
    for shipment in shipments:
        if shipment.shipment_status == "archived":
            continue
        container_number = _clean_container(shipment.container_number)
        normalized_bl = _normalize_bl_number(shipment.bl_number)
        if normalized_bl and (container_number, normalized_bl) in archived_group_keys:
            continue
        if not normalized_bl and container_number in archived_blank_containers:
            continue
        if not _normalize_bl_number(shipment.bl_number) and _shipment_identity_key(shipment) in suppressed_orphan_keys:
            continue
        row = _shipment_to_dict(shipment)
        normalized_bl = _normalize_bl_number(row.get("bl_number"))
        group_key = f"BL:{normalized_bl}" if normalized_bl else f"SHIP:{row['id']}"
        grouped.setdefault(group_key, []).append(row)

    rows: list[dict[str, Any]] = []
    for group_key, entries in grouped.items():
        sorted_entries = sorted(
            entries,
            key=lambda item: (
                MOVEMENT_PRIORITY.get(item.get("movement_category") or "Hi Seas", 0),
                bool(_clean_text(item.get("train_no"))),
                _date_sort_value(item.get("latest_time", "")),
                item.get("id", 0),
            ),
            reverse=True,
        )
        lead = sorted_entries[0]
        refresh_entries = sorted(
            sorted_entries,
            key=lambda item: (
                _date_sort_value(item.get("last_refresh_at", "")),
                bool(_clean_text(item.get("last_refresh_status", ""))),
                item.get("id", 0),
            ),
            reverse=True,
        )
        refresh_lead = refresh_entries[0]
        bl_number = _first_non_empty([item.get("bl_number", "") for item in sorted_entries])
        container_numbers = sorted({item.get("container_number", "") for item in sorted_entries if item.get("container_number")})
        shipment_status = max(
            (item.get("shipment_status") or "archived" for item in sorted_entries),
            key=lambda status: SHIPMENT_STATUS_PRIORITY.get(status, 0),
        )
        tracking_sources: list[str] = []
        for item in sorted_entries:
            for source in re.split(r"[,+]", _clean_text(item.get("tracking_source", "")).lower()):
                source = source.strip()
                if source and source not in tracking_sources:
                    tracking_sources.append(source)
        documents = _documents_for_bl(bl_number)
        lead_diagnostics = dict(lead.get("movement_diagnostics") or {})
        if lead_diagnostics:
            lead_diagnostics["group_scope_summary"] = (
                f"The dashboard is showing the strongest live movement across {len(container_numbers)} container"
                f"{'' if len(container_numbers) == 1 else 's'} in this BL group."
            )
            if any(bool(item.get("action_required")) for item in sorted_entries):
                lead_diagnostics["action_summary"] = _first_non_empty(
                    [item.get("action_required_reason", "") for item in sorted_entries]
                )
        rows.append(
            {
                "group_key": group_key,
                "id": lead.get("id"),
                "customer_name": _first_non_empty([item.get("customer_name", "") for item in sorted_entries]),
                "primary_container_number": lead.get("container_number", ""),
                "container_number": lead.get("container_number", ""),
                "container_numbers": container_numbers,
                "container_count": len(container_numbers),
                "bl_number": bl_number,
                "shipment_status": shipment_status,
                "movement_category": lead.get("movement_category") or "Hi Seas",
                "latest_location": _first_non_empty([item.get("latest_location", "") for item in sorted_entries]),
                "latest_time": _first_non_empty([item.get("latest_time", "") for item in sorted_entries]),
                "movement_since_date": _earliest_non_empty_date([item.get("movement_since_date", "") for item in sorted_entries]),
                "port_arrival_date": _earliest_non_empty_date([item.get("port_arrival_date", "") for item in sorted_entries]),
                "birgunj_arrival_date": _earliest_non_empty_date([item.get("birgunj_arrival_date", "") for item in sorted_entries]),
                "pristine_booking_date": _earliest_non_empty_date([item.get("pristine_booking_date", "") for item in sorted_entries]),
                "train_no": _first_non_empty([item.get("train_no", "") for item in sorted_entries]),
                "departure": _first_non_empty([item.get("departure", "") for item in sorted_entries]),
                "wagon_loaded_date": _earliest_non_empty_date([item.get("wagon_loaded_date", "") for item in sorted_entries]),
                "tracking_source": ",".join(tracking_sources),
                "source_type": _first_non_empty([item.get("source_type", "") for item in sorted_entries]),
                "source_label": _first_non_empty([item.get("source_label", "") for item in sorted_entries]),
                "source_batch_id": refresh_lead.get("source_batch_id") or lead.get("source_batch_id") or 0,
                "last_refresh_at": _first_non_empty([item.get("last_refresh_at", "") for item in refresh_entries]),
                "last_refresh_status": _first_non_empty([item.get("last_refresh_status", "") for item in refresh_entries]),
                "last_refresh_error": _first_non_empty([item.get("last_refresh_error", "") for item in refresh_entries]),
                "clearance_doc_number": _first_non_empty([item.get("clearance_doc_number", "") for item in sorted_entries]),
                "do_date": _first_non_empty([item.get("do_date", "") for item in sorted_entries]),
                "document_status": _first_non_empty([item.get("document_status", "") for item in sorted_entries]),
                "original_docs_received_date": _first_non_empty([item.get("original_docs_received_date", "") for item in sorted_entries]),
                "action_required": any(bool(item.get("action_required")) for item in sorted_entries),
                "action_required_reason": _first_non_empty([item.get("action_required_reason", "") for item in sorted_entries]),
                "movement_diagnostics": lead_diagnostics,
                "documents": documents,
                "documents_complete": bool(bl_number) and all(documents.get(doc_type) for doc_type in VALID_DOCUMENT_TYPES),
            }
        )

    rows.sort(
        key=lambda item: (
            SHIPMENT_STATUS_PRIORITY.get(item.get("shipment_status") or "archived", 0),
            MOVEMENT_PRIORITY.get(item.get("movement_category") or "Hi Seas", 0),
            _date_sort_value(item.get("latest_time", "")),
            item.get("customer_name", ""),
        ),
        reverse=True,
    )
    return rows


def _format_customer_count_items(counter: dict[str, int]) -> list[dict[str, Any]]:
    return [
        {"customer_name": name, "shipment_count": count, "container_count": count}
        for name, count in sorted(counter.items(), key=lambda item: (-item[1], item[0]))
    ]


def _dashboard_identifiers_from_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    today_text = datetime.now().strftime("%d-%m-%Y")
    rolling_week_cutoff = (datetime.now() - timedelta(days=7)).date()
    total_at_icd_birgunj = 0
    today_arrivals = 0
    approaching_birgunj = 0
    railed_out_this_week = 0
    today_arrival_customers: dict[str, int] = {}
    approaching_birgunj_customers: dict[str, int] = {}
    railed_out_this_week_customers: dict[str, int] = {}

    for row in rows:
        movement_category = _clean_text(row.get("movement_category")) or "Hi Seas"
        customer_name = _format_customer_name(row.get("customer_name", ""))
        arrived = movement_category == "Arrived Birgunj"
        rail_start_date = (
            _parse_date(row.get("departure", ""))
            or _parse_date(row.get("wagon_loaded_date", ""))
            or (_parse_date(row.get("movement_since_date", "")) if movement_category == "On Rail" else None)
        )
        if arrived:
            total_at_icd_birgunj += 1
            arrival_date = _clean_text(row.get("birgunj_arrival_date", "")) or _clean_text(row.get("latest_time", ""))
            if _format_to_dd_mm_yyyy(arrival_date) == today_text:
                today_arrivals += 1
                if customer_name:
                    today_arrival_customers[customer_name] = today_arrival_customers.get(customer_name, 0) + 1
            continue

        location_text = _clean_text(row.get("latest_location", ""))
        if not location_text:
            continue
        distance_info = _resolve_location_distance(location_text)
        distance_km = distance_info.get("distance_km")
        if distance_info.get("found") and isinstance(distance_km, (int, float)) and float(distance_km) < 40:
            approaching_birgunj += 1
            if customer_name:
                approaching_birgunj_customers[customer_name] = approaching_birgunj_customers.get(customer_name, 0) + 1

        if rail_start_date and rail_start_date.date() >= rolling_week_cutoff:
            railed_out_this_week += 1
            if customer_name:
                railed_out_this_week_customers[customer_name] = railed_out_this_week_customers.get(customer_name, 0) + 1

    return {
        "total_at_icd_birgunj": total_at_icd_birgunj,
        "today_arrivals": today_arrivals,
        "approaching_birgunj": approaching_birgunj,
        "railed_out_this_week": railed_out_this_week,
        "today_arrival_customers": _format_customer_count_items(today_arrival_customers),
        "approaching_birgunj_customers": _format_customer_count_items(approaching_birgunj_customers),
        "railed_out_this_week_customers": _format_customer_count_items(railed_out_this_week_customers),
    }


def _dashboard_identifiers(shipments: list[Shipment]) -> dict[str, Any]:
    return _dashboard_identifiers_from_rows(_group_dashboard_rows(shipments))


def _shipment_count_summary(shipments: list[Shipment], dashboard_rows: list[dict[str, Any]] | None = None) -> dict[str, int]:
    rows = dashboard_rows if dashboard_rows is not None else _group_dashboard_rows(shipments)
    completed_groups = {
        f"BL:{_normalize_bl_number(shipment.bl_number)}"
        if _normalize_bl_number(shipment.bl_number)
        else f"SHIP:{shipment.id}"
        for shipment in shipments
        if shipment.shipment_status == "completed"
    }
    archived_groups = {
        f"BL:{_normalize_bl_number(shipment.bl_number)}"
        if _normalize_bl_number(shipment.bl_number)
        else f"SHIP:{shipment.id}"
        for shipment in shipments
        if shipment.shipment_status == "archived"
    }
    return {
        "active": len(rows),
        "completed": len(completed_groups),
        "archived": len(archived_groups),
        "total": len(rows) + len(completed_groups) + len(archived_groups),
    }


@router.get("")
def list_shipments(db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    return [_shipment_to_dict(shipment) for shipment in _get_shipments(db, current_user)]


@router.get("/sources")
def list_sources(db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    sources = list(
        db.execute(
            select(ShipmentSource)
            .where(ShipmentSource.user_id == current_user.id)
            .order_by(ShipmentSource.updated_at.desc(), ShipmentSource.id.desc())
        ).scalars()
    )
    return [
        {
            "id": source.id,
            "source_type": source.source_type,
            "source_label": source.source_label,
            "source_reference": source.source_reference,
            "status": source.status,
            "last_sync_at": source.last_sync_at.isoformat() if source.last_sync_at else "",
        }
        for source in sources
    ]


@router.get("/source-mappings")
def list_source_mappings(db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    return _serialize_source_mappings(db, current_user)


@router.get("/source-connections")
def list_source_connections(db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    return _serialize_source_connections(db, current_user)


@router.post("/source-connections/google-sheets")
def create_google_sheets_connection(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    source_url = _clean_text(payload.get("source_url"))
    worksheet_name = _clean_text(payload.get("worksheet_name"))
    connection_label = _clean_text(payload.get("connection_label")) or "Google Sheet"
    mapping_profile_id = int(payload.get("mapping_profile_id") or 0)
    parsed = _parse_google_sheet_reference(source_url)

    source = _get_or_create_source(
        db,
        current_user,
        "google_sheets",
        connection_label,
        source_reference=parsed["sheet_id"],
        metadata={
            "sheet_id": parsed["sheet_id"],
            "gid": parsed["gid"],
            "worksheet_name": worksheet_name,
        },
    )

    existing = db.execute(
        select(SourceConnection).where(
            SourceConnection.user_id == current_user.id,
            SourceConnection.provider == "google_sheets",
            SourceConnection.source_url == source_url,
            SourceConnection.worksheet_name == worksheet_name,
        )
    ).scalar_one_or_none()

    if existing:
        existing.connection_label = connection_label
        existing.mapping_profile_id = mapping_profile_id
        existing.status = "connected"
        existing.config_json = _json_dumps(parsed)
        existing.updated_at = datetime.utcnow()
        connection = existing
    else:
        connection = SourceConnection(
            user_id=current_user.id,
            source_id=source.id,
            provider="google_sheets",
            connection_label=connection_label,
            source_url=source_url,
            worksheet_name=worksheet_name,
            mapping_profile_id=mapping_profile_id,
            status="connected",
            config_json=_json_dumps(parsed),
        )
        db.add(connection)
        db.flush()

    _log_audit_event(
        db,
        current_user.id,
        "source_connection_saved",
        shipment_status="active",
        details={
            "provider": "google_sheets",
            "connection_label": connection_label,
            "worksheet_name": worksheet_name,
            "sheet_id": parsed["sheet_id"],
            "mapping_profile_id": mapping_profile_id,
        },
    )
    db.commit()
    return {
        "id": connection.id,
        "provider": connection.provider,
        "connection_label": connection.connection_label,
        "source_url": connection.source_url,
        "worksheet_name": connection.worksheet_name,
        "mapping_profile_id": connection.mapping_profile_id,
        "status": connection.status,
        "config": parsed,
    }


@router.post("/source-connections/google-sheets/preview")
def preview_google_sheets_source(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    source_url = _clean_text(payload.get("source_url"))
    worksheet_name = _clean_text(payload.get("worksheet_name"))
    workbook_bytes, parsed = _download_google_sheet_workbook(source_url)
    workbook = load_workbook(filename=BytesIO(workbook_bytes), data_only=True)
    sheet_title = _fetch_google_sheet_title(source_url) or _sheet_title_from_workbook(workbook, parsed)
    available_sheets = list(workbook.sheetnames)
    if not available_sheets:
        raise HTTPException(status_code=400, detail="Google Sheet does not contain any readable worksheets")

    token = f"{uuid4().hex}.xlsx"
    file_path = TEMP_IMPORT_DIR / token
    file_path.write_bytes(workbook_bytes)

    selected_sheet = worksheet_name or available_sheets[0]
    sheet_name, header_row, headers, preview_rows, _rows, _available_sheets = _read_sheet(file_path, selected_sheet)
    remembered_mapping = _load_recent_mapping_for_sheet(db, current_user, sheet_name, "google_sheets")
    remembered_profile = db.execute(
        select(SourceMappingProfile)
        .where(
            SourceMappingProfile.user_id == current_user.id,
            SourceMappingProfile.source_type == "google_sheets",
            SourceMappingProfile.source_sheet == _clean_text(sheet_name),
        )
        .order_by(SourceMappingProfile.updated_at.desc(), SourceMappingProfile.id.desc())
        .limit(1)
    ).scalar_one_or_none()

    session = UploadSession(
        user_id=current_user.id,
        original_filename=f"{sheet_title}.xlsx",
        stored_path=str(file_path),
        detected_sheet=sheet_name,
        detected_header_row=header_row,
        available_columns_json=headers,
        preview_rows_json=preview_rows,
        status="preview_ready",
    )
    db.add(session)
    db.commit()

    return {
        "temp_file_token": token,
        "upload_session_id": session.id,
        "sheet_title": sheet_title,
        "sheet_id": parsed["sheet_id"],
        "source_url": source_url,
        "available_sheets": available_sheets,
        "sheet_name": sheet_name,
        "header_row": header_row,
        "available_columns": headers,
        "preview_rows": preview_rows,
        "remembered_mapping": remembered_mapping,
        "remembered_profile": {
            "id": remembered_profile.id,
            "profile_label": remembered_profile.profile_label,
            "source_sheet": remembered_profile.source_sheet,
        } if remembered_profile else None,
        "source_context": {
            "source_type": "google_sheets",
            "source_url": source_url,
            "sheet_id": parsed["sheet_id"],
            "sheet_title": sheet_title,
            "worksheet_name": sheet_name,
        },
    }


@router.get("/source-batches")
def list_source_batches(
    limit: int = 20,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_portal_user),
):
    _ensure_user_scope_ready(db, current_user)
    return _serialize_source_batches(db, current_user, limit=limit)


@router.get("/source-batches/{batch_id}")
def get_source_batch_detail(
    batch_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_portal_user),
):
    _ensure_user_scope_ready(db, current_user)
    batch = db.get(ShipmentBatch, batch_id)
    if not batch or batch.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Source batch not found")
    source = db.get(ShipmentSource, batch.source_id) if batch.source_id else None
    raw_rows = list(
        db.execute(
            select(RawSourceRow)
            .where(
                RawSourceRow.user_id == current_user.id,
                RawSourceRow.batch_id == batch.id,
            )
            .order_by(RawSourceRow.source_row_number.asc(), RawSourceRow.id.asc())
            .limit(50)
        ).scalars()
    )
    return {
        "id": batch.id,
        "batch_label": batch.batch_label,
        "source_sheet": batch.source_sheet,
        "header_row": batch.header_row,
        "status": batch.status,
        "imported_count": batch.imported_count,
        "duplicate_count": batch.duplicate_count,
        "invalid_count": batch.invalid_count,
        "skipped_blank_count": batch.skipped_blank_count,
        "created_at": batch.created_at.isoformat(),
        "source": {
            "id": source.id if source else 0,
            "source_type": source.source_type if source else "",
            "source_label": source.source_label if source else "",
            "source_reference": source.source_reference if source else "",
            "last_sync_at": source.last_sync_at.isoformat() if source and source.last_sync_at else "",
        },
        "rows": [
            {
                "id": row.id,
                "source_row_number": row.source_row_number,
                "row_status": row.row_status,
                "row_error": row.row_error,
                "raw_row": json.loads(row.raw_row_json or "{}"),
                "mapped_row": json.loads(row.mapped_row_json or "{}"),
            }
            for row in raw_rows
        ],
    }


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    shipments = _get_shipments(db, current_user)
    rows = _group_dashboard_rows(shipments)
    return {
        "rows": rows,
        "identifiers": _dashboard_identifiers_from_rows(rows),
    }


@router.get("/bootstrap")
def dashboard_bootstrap(
    include_sources: bool = True,
    include_shipments: bool = True,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_portal_user),
):
    _ensure_user_scope_ready(db, current_user)
    shipments = _get_shipments(db, current_user)
    dashboard_rows = _group_dashboard_rows(shipments)
    response = {
        "dashboard": {
            "rows": dashboard_rows,
            "identifiers": _dashboard_identifiers_from_rows(dashboard_rows),
            "shipment_counts": _shipment_count_summary(shipments, dashboard_rows),
        },
    }
    if include_shipments:
        response["shipments"] = [_shipment_to_dict(shipment) for shipment in shipments]
    if include_sources:
        response.update(
            {
                "source_batches": _serialize_source_batches(db, current_user, limit=8),
                "source_mappings": _serialize_source_mappings(db, current_user),
                "source_connections": _serialize_source_connections(db, current_user),
            }
        )
    return response


@router.get("/dashboard-export")
def export_dashboard(db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    rows = _group_dashboard_rows(_get_shipments(db, current_user))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Live Dashboard"

    headers = [
        "Customer",
        "BL Number",
        "Containers",
        "Shipment Status",
        "Movement",
        "Latest Location",
        "Movement Since",
        "Train No",
        "Departure",
        "DO Date",
        "Document Status",
        "Original Docs Received",
        "Clearance Doc",
        "Action Needed",
    ]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D0D7E2"),
        right=Side(style="thin", color="D0D7E2"),
        top=Side(style="thin", color="D0D7E2"),
        bottom=Side(style="thin", color="D0D7E2"),
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row in rows:
        sheet.append(
            [
                row.get("customer_name", ""),
                row.get("bl_number", ""),
                ", ".join(row.get("container_numbers") or []),
                row.get("shipment_status", ""),
                row.get("movement_category", ""),
                row.get("latest_location", ""),
                row.get("movement_since_date") or row.get("latest_time", ""),
                row.get("train_no", ""),
                row.get("departure", ""),
                row.get("do_date", ""),
                row.get("document_status", ""),
                row.get("original_docs_received_date", ""),
                row.get("clearance_doc_number", ""),
                "Yes" if row.get("action_required") else "",
            ]
        )

    widths = [28, 18, 40, 16, 18, 34, 16, 12, 16, 16, 18, 20, 18, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="tracking-dashboard.xlsx"'},
    )


@router.get("/customers")
def customer_suggestions(q: str = "", db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    query = _clean_text(q).upper()
    items = sorted({_format_customer_name(shipment.customer_name) for shipment in _get_shipments(db, current_user) if shipment.customer_name})
    if query:
        items = [item for item in items if query in item.upper()]
    return {"items": items[:12]}


@router.post("/location-distances")
def location_distances(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    locations = payload.get("locations") or []
    if not isinstance(locations, list):
        raise HTTPException(status_code=400, detail="locations must be a list")
    results: dict[str, Any] = {}
    for location in locations:
        location_text = _clean_text(location)
        if not location_text:
            continue
        results[location_text] = _resolve_location_distance(location_text)
    return {"items": results}


@router.post("/add")
def add_shipment(payload: ShipmentCreateRequest, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    container_numbers, invalid_containers = _parse_container_numbers(
        payload.container_numbers,
        payload.container_number,
    )
    if invalid_containers:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid container format: {', '.join(invalid_containers)}. Use 4 letters followed by 7 digits.",
        )
    if not container_numbers:
        raise HTTPException(status_code=400, detail="At least one container number is required")
    canonical_customer = _sync_customer_directory(db, payload.customer_name)
    normalized_bl = _normalize_bl_number(payload.bl_number)
    existing_records = list(
        db.execute(
            select(Shipment).where(
                Shipment.user_id == current_user.id,
                Shipment.shipment_status != "archived",
                Shipment.bl_number == normalized_bl,
                Shipment.customer_name == canonical_customer,
                Shipment.container_number.in_(container_numbers),
            )
        ).scalars()
    )
    existing_containers = {_clean_container(shipment.container_number) for shipment in existing_records}
    duplicate_containers = [container_number for container_number in container_numbers if container_number in existing_containers]
    if duplicate_containers:
        raise HTTPException(
            status_code=400,
            detail=f"Duplicate shipment for: {', '.join(duplicate_containers)}",
        )

    created_shipments: list[Shipment] = []
    for container_number in container_numbers:
        shipment = Shipment(
            user_id=current_user.id,
            customer_name=canonical_customer,
            container_number=container_number,
            bl_number=normalized_bl,
            shipment_status="active",
            movement_category="Hi Seas",
            source_type="manual",
            source_label="Manual Entry",
        )
        db.add(shipment)
        created_shipments.append(shipment)
    _log_audit_event(
        db,
        current_user.id,
        "shipment_added",
        bl_number=normalized_bl,
        container_number=container_numbers[0],
        shipment_status="active",
        details={"customer_name": canonical_customer, "container_numbers": container_numbers, "created_count": len(container_numbers)},
    )
    db.commit()
    for shipment in created_shipments:
        db.refresh(shipment)
    return {
        "created_count": len(created_shipments),
        "items": [_shipment_to_dict(shipment) for shipment in created_shipments],
    }


@router.patch("/actions/group/edit")
def update_group_details(
    payload: ShipmentGroupUpdateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_portal_user),
):
    _ensure_user_scope_ready(db, current_user)
    container_numbers, invalid_containers = _parse_container_numbers(payload.container_numbers)
    if invalid_containers:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid container format: {', '.join(invalid_containers)}. Use 4 letters followed by 7 digits.",
        )
    if not container_numbers:
        raise HTTPException(status_code=400, detail="At least one valid container number is required")

    target_shipments = _resolve_group_shipments(
        db,
        current_user,
        bl_number=payload.current_bl_number,
        container_numbers=payload.current_container_numbers,
        include_archived=True,
    )
    if not target_shipments:
        raise HTTPException(status_code=404, detail="Shipment group not found")

    canonical_customer = _sync_customer_directory(db, payload.customer_name)
    normalized_bl = _normalize_bl_number(payload.bl_number)
    normalized_clearance_doc_number = _clean_text(payload.clearance_doc_number)
    normalized_do_date = _normalize_manual_date(payload.do_date)
    normalized_document_status = _normalize_document_status(payload.document_status)
    normalized_original_docs_received_date = _normalize_manual_date(payload.original_docs_received_date)
    if normalized_document_status != "Original":
        normalized_original_docs_received_date = ""
    if normalized_document_status == "Original" and not normalized_original_docs_received_date:
        raise HTTPException(status_code=400, detail="Original document date is required when document status is Original")
    previous_bl = _first_non_empty([shipment.bl_number for shipment in target_shipments])
    active_status = max(
        (shipment.shipment_status or "active" for shipment in target_shipments),
        key=lambda status: SHIPMENT_STATUS_PRIORITY.get(status, 0),
    )

    existing_elsewhere = list(
        db.execute(
            _user_shipment_select(current_user).where(
                Shipment.container_number.in_(container_numbers),
                ~Shipment.id.in_([shipment.id for shipment in target_shipments]),
            )
        ).scalars()
    )
    if existing_elsewhere:
        taken = sorted({_clean_container(shipment.container_number) for shipment in existing_elsewhere if _clean_container(shipment.container_number)})
        raise HTTPException(
            status_code=400,
            detail=f"These containers already exist in another shipment: {', '.join(taken)}",
        )

    existing_by_container = {
        _clean_container(shipment.container_number): shipment
        for shipment in target_shipments
        if _clean_container(shipment.container_number)
    }
    next_containers = set(container_numbers)

    for shipment in target_shipments:
        if _clean_container(shipment.container_number) not in next_containers:
            db.delete(shipment)

    updated_shipments: list[Shipment] = []
    for container_number in container_numbers:
        shipment = existing_by_container.get(container_number)
        if shipment is None:
            shipment = Shipment(
                user_id=current_user.id,
                customer_name=canonical_customer,
                container_number=container_number,
                bl_number=normalized_bl,
                shipment_status=active_status,
                movement_category="Hi Seas",
                source_type="manual",
                source_label="Manual Entry",
            )
            db.add(shipment)
        else:
            shipment.customer_name = canonical_customer
            shipment.bl_number = normalized_bl
        shipment.clearance_doc_number = normalized_clearance_doc_number
        shipment.do_date = normalized_do_date
        shipment.document_status = normalized_document_status
        shipment.original_docs_received_date = normalized_original_docs_received_date
        updated_shipments.append(shipment)

    if previous_bl and normalized_bl and previous_bl != normalized_bl:
        _move_documents_between_bls(previous_bl, normalized_bl)

    _log_audit_event(
        db,
        current_user.id,
        "shipment_group_edited",
        bl_number=normalized_bl or previous_bl,
        container_number=container_numbers[0],
        shipment_status=active_status,
        details={
            "customer_name": canonical_customer,
            "container_numbers": container_numbers,
            "previous_bl_number": previous_bl,
            "next_bl_number": normalized_bl,
            "clearance_doc_number": normalized_clearance_doc_number,
            "do_date": normalized_do_date,
            "document_status": normalized_document_status,
            "original_docs_received_date": normalized_original_docs_received_date,
        },
    )
    db.commit()

    refreshed_shipments = _resolve_group_shipments(
        db,
        current_user,
        bl_number=normalized_bl or previous_bl,
        container_numbers=container_numbers,
        include_archived=True,
    )
    rows = _group_dashboard_rows(refreshed_shipments)
    return {
        "updated_count": len(container_numbers),
        "items": [_shipment_to_dict(shipment) for shipment in refreshed_shipments],
        "row": rows[0] if rows else None,
    }


@router.patch("/{shipment_id}/status")
def update_status(shipment_id: int, payload: ShipmentStatusUpdateRequest, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    next_status = _clean_text(payload.shipment_status).lower()
    if next_status not in {"active", "completed", "archived"}:
        raise HTTPException(status_code=400, detail="Invalid shipment status")
    shipment = db.get(Shipment, shipment_id)
    if not shipment or shipment.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Not found")
    shipment.shipment_status = next_status
    db.commit()
    db.refresh(shipment)
    return _shipment_to_dict(shipment)


@router.patch("/actions/group/status")
def update_group_status(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    next_status = _clean_text(payload.get("shipment_status")).lower()
    if next_status not in {"active", "completed", "archived"}:
        raise HTTPException(status_code=400, detail="Invalid shipment status")
    clearance_doc_number = _clean_text(payload.get("clearance_doc_number"))
    shipments, effective_doc_number = _apply_group_status_transition(
        db,
        current_user,
        bl_number=payload.get("bl_number"),
        container_numbers=payload.get("container_numbers") or [],
        next_status=next_status,
        clearance_doc_number=clearance_doc_number,
    )
    db.commit()
    return {"updated": True, "count": len(shipments), "shipment_status": next_status, "clearance_doc_number": effective_doc_number}


@router.patch("/actions/bulk/status")
def update_bulk_group_status(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    next_status = _clean_text(payload.get("shipment_status")).lower()
    if next_status not in {"active", "completed", "archived"}:
        raise HTTPException(status_code=400, detail="Invalid shipment status")
    groups = payload.get("groups") or []
    if not isinstance(groups, list) or not groups:
        raise HTTPException(status_code=400, detail="At least one shipment group is required")
    clearance_map = payload.get("clearance_doc_numbers") or {}
    if not isinstance(clearance_map, dict):
        clearance_map = {}

    updated_group_count = 0
    updated_row_ids: set[int] = set()
    for group in groups:
        if not isinstance(group, dict):
            continue
        group_key = _clean_text(group.get("group_key"))
        shipments, _effective_doc_number = _apply_group_status_transition(
            db,
            current_user,
            bl_number=group.get("bl_number"),
            container_numbers=group.get("container_numbers") or [],
            next_status=next_status,
            clearance_doc_number=_clean_text(clearance_map.get(group_key)),
        )
        updated_group_count += 1
        updated_row_ids.update(shipment.id for shipment in shipments)

    db.commit()
    return {
        "updated": True,
        "shipment_status": next_status,
        "group_count": updated_group_count,
        "row_count": len(updated_row_ids),
    }


@router.delete("/{shipment_id}")
def delete_shipment(shipment_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    shipment = db.get(Shipment, shipment_id)
    if not shipment or shipment.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Not found")
    db.delete(shipment)
    db.commit()
    return {"deleted": True}


@router.delete("/group")
def delete_shipment_group(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    shipments = _delete_group_shipments_internal(
        db,
        current_user,
        bl_number=payload.get("bl_number"),
        container_numbers=payload.get("container_numbers") or [],
    )
    db.commit()
    return {"deleted": True, "count": len(shipments)}


@router.post("/group/delete")
def delete_shipment_group_post(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    return delete_shipment_group(payload, db, current_user)


@router.post("/group/delete-bulk")
def delete_shipment_group_bulk(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    groups = payload.get("groups") or []
    if not isinstance(groups, list) or not groups:
        raise HTTPException(status_code=400, detail="At least one shipment group is required")

    deleted_group_count = 0
    deleted_row_count = 0
    for group in groups:
        if not isinstance(group, dict):
            continue
        shipments = _delete_group_shipments_internal(
            db,
            current_user,
            bl_number=group.get("bl_number"),
            container_numbers=group.get("container_numbers") or [],
        )
        deleted_group_count += 1
        deleted_row_count += len(shipments)

    db.commit()
    return {"deleted": True, "group_count": deleted_group_count, "count": deleted_row_count}

@router.post("/import-preview")
async def import_preview(file: UploadFile = File(...), db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    filename = file.filename or ""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise HTTPException(status_code=400, detail="Only .xlsx and .xlsm files are supported")
    file_bytes = await file.read()
    max_upload_bytes = max(1, int(settings.max_upload_mb)) * 1024 * 1024
    if len(file_bytes) > max_upload_bytes:
        raise HTTPException(
            status_code=413,
            detail=f"Workbook exceeds the {settings.max_upload_mb} MB upload limit.",
        )
    token = f"{uuid4().hex}{suffix}"
    file_path = TEMP_IMPORT_DIR / token
    file_path.write_bytes(file_bytes)
    sheet_name, header_row, headers, preview_rows, _rows, available_sheets = _read_sheet(file_path)
    session = UploadSession(
        user_id=current_user.id,
        original_filename=filename,
        stored_path=str(file_path),
        detected_sheet=sheet_name,
        detected_header_row=header_row,
        available_columns_json=headers,
        preview_rows_json=preview_rows,
        status="preview_ready",
    )
    db.add(session)
    db.commit()
    remembered_mapping = _load_recent_mapping_for_sheet(db, current_user, sheet_name)
    remembered_profile = db.execute(
        select(SourceMappingProfile)
        .where(
            SourceMappingProfile.user_id == current_user.id,
            SourceMappingProfile.source_type == "excel_upload",
            SourceMappingProfile.source_sheet == _clean_text(sheet_name),
        )
        .order_by(SourceMappingProfile.updated_at.desc(), SourceMappingProfile.id.desc())
        .limit(1)
    ).scalar_one_or_none()
    return {
        "temp_file_token": token,
        "upload_session_id": session.id,
        "original_filename": filename,
        "sheet_name": sheet_name,
        "header_row": header_row,
        "available_columns": headers,
        "available_sheets": available_sheets,
        "preview_rows": preview_rows,
        "remembered_mapping": remembered_mapping,
        "remembered_profile": {
            "id": remembered_profile.id,
            "profile_label": remembered_profile.profile_label,
            "source_sheet": remembered_profile.source_sheet,
        } if remembered_profile else None,
    }


@router.post("/import-confirm")
async def import_confirm(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    token = _clean_text(payload.get("temp_file_token"))
    mapping_json = payload.get("mapping_json") or {}
    source_context = payload.get("source_context") or {}
    if not token:
        raise HTTPException(status_code=400, detail="temp_file_token is required")
    file_path = TEMP_IMPORT_DIR / token
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Temporary import file not found")
    upload_session = db.execute(
        select(UploadSession).where(
            UploadSession.user_id == current_user.id,
            UploadSession.stored_path == str(file_path),
        )
    ).scalar_one_or_none()
    customer_col = _clean_text(mapping_json.get("customer_name"))
    container_col = _clean_text(mapping_json.get("container_number"))
    bl_col = _clean_text(mapping_json.get("bl_number"))
    row_overrides = payload.get("row_overrides") or []
    if not container_col:
        raise HTTPException(status_code=400, detail="Container Number mapping is required")
    selected_sheet = upload_session.detected_sheet if upload_session else None
    sheet_name, header_row, headers, _preview_rows, rows, _available_sheets = _read_sheet(file_path, selected_sheet)
    header_index = header_row - 1
    normalized_rows = _normalize_import_row_objects(
        headers,
        rows,
        header_index,
        [customer_col, bl_col],
    )
    normalized_rows = _apply_import_row_overrides(
        normalized_rows,
        row_overrides,
        customer_col,
        container_col,
        bl_col,
    )
    review = _summarize_import_rows(normalized_rows, customer_col, container_col, bl_col)
    if review["invalid_count"]:
        raise HTTPException(
            status_code=409,
            detail={
                "message": "Action required before import",
                "invalid_rows": review["invalid_rows"],
                "invalid_count": review["invalid_count"],
                "valid_count": review["valid_count"],
                "skipped_blank_count": review["skipped_blank_count"],
            },
        )
    existing_keys = {
        (
            shipment.container_number,
            _normalize_bl_number(shipment.bl_number),
            _format_customer_name(shipment.customer_name),
        )
        for shipment in db.execute(
            _user_shipment_select(current_user).where(Shipment.shipment_status != "archived")
        ).scalars()
    }
    imported_count = 0
    duplicate_count = 0
    skipped_blank_count = 0
    skipped_invalid_count = 0
    user_id = current_user.id
    source_type = _clean_text(source_context.get("source_type")) or "excel_upload"
    source_url = _clean_text(source_context.get("source_url"))
    sheet_id = _clean_text(source_context.get("sheet_id"))
    sheet_title = _clean_text(source_context.get("sheet_title"))
    source_label = "Excel Upload"
    source_reference = upload_session.original_filename if upload_session else Path(file_path).name
    source_metadata = {
        "original_filename": upload_session.original_filename if upload_session else Path(file_path).name,
        "sheet_name": sheet_name,
    }
    if source_type == "google_sheets":
        source_label = sheet_title or "Google Sheet"
        source_reference = sheet_id or source_url or source_reference
        source_metadata = {
            "source_url": source_url,
            "sheet_id": sheet_id,
            "worksheet_name": sheet_name,
        }
    source = None
    mapping_profile = None
    batch = None
    raw_rows_by_source_number: dict[int, RawSourceRow] = {}
    source_tracking_warning = ""
    try:
        source = _get_or_create_source(
            db,
            current_user,
            source_type,
            source_label,
            source_reference=source_reference,
            metadata=source_metadata,
        )
        mapping_profile = _upsert_mapping_profile(
            db,
            current_user,
            source_type=source_type,
            source_sheet=sheet_name,
            profile_label=(f"Google Sheets - {sheet_name}" if source_type == "google_sheets" else f"Excel - {sheet_name}"),
            mapping_json=mapping_json,
        )
        batch = _create_import_batch(
            db,
            current_user,
            source,
            batch_label=upload_session.original_filename if upload_session else Path(file_path).name,
            source_sheet=sheet_name,
            header_row=header_row,
            mapping_json=mapping_json,
            imported_count=0,
            duplicate_count=0,
            invalid_count=review["invalid_count"],
            skipped_blank_count=review["skipped_blank_count"],
        )
        raw_rows_by_source_number = _record_raw_import_rows(
            db,
            current_user,
            source,
            batch,
            normalized_rows,
            customer_col,
            container_col,
            bl_col,
        )
    except Exception:
        logger.exception(
            "Import provenance save failed for user %s and source %s",
            current_user.id,
            source_type,
        )
        db.rollback()
        source = None
        mapping_profile = None
        batch = None
        raw_rows_by_source_number = {}
        source_tracking_warning = (
            "This import was added without source memory. Shipment intake still succeeded, but batch provenance could not be saved."
        )
    for row_obj in normalized_rows:
        source_row_number = int(row_obj.get("__source_row_number") or 0)
        raw_row = raw_rows_by_source_number.get(source_row_number)
        valid_containers, invalid_containers, raw_container_value = _containers_from_import_row(row_obj, container_col)
        if not valid_containers and not invalid_containers:
            skipped_blank_count += 1
            if raw_row:
                raw_row.row_status = "blank"
            continue
        if invalid_containers:
            skipped_invalid_count += len(invalid_containers)
            if raw_row:
                raw_row.row_status = "invalid"
                raw_row.row_error = f"{' '.join(invalid_containers)} must use 4 letters followed by 7 digits."
            continue
        canonical_customer = _sync_customer_directory(db, row_obj.get(customer_col))
        bl_number = _normalize_bl_number(row_obj.get(bl_col))
        row_imported = 0
        row_duplicates = 0
        for container_number in valid_containers:
            shipment_key = (container_number, bl_number, canonical_customer)
            if shipment_key in existing_keys:
                duplicate_count += 1
                row_duplicates += 1
                continue
            shipment = Shipment(
                user_id=user_id,
                customer_name=canonical_customer,
                container_number=container_number,
                bl_number=bl_number,
                shipment_status="active",
                movement_category="Hi Seas",
                source_type=source_type,
                source_label=source_label,
                source_batch_id=batch.id if batch else 0,
                raw_source_row_id=raw_row.id if raw_row else 0,
            )
            db.add(shipment)
            existing_keys.add(shipment_key)
            imported_count += 1
            row_imported += 1
        if raw_row:
            if row_imported:
                raw_row.row_status = "imported"
                raw_row.row_error = ""
            elif row_duplicates and row_duplicates == len(valid_containers):
                raw_row.row_status = "duplicate"
    if batch:
        batch.imported_count = imported_count
        batch.duplicate_count = duplicate_count
        batch.invalid_count = skipped_invalid_count
        batch.skipped_blank_count = skipped_blank_count
    if source:
        source.last_sync_at = datetime.utcnow()
    if source_type == "google_sheets" and source_url and source and mapping_profile:
        parsed = _parse_google_sheet_reference(source_url)
        existing_connection = db.execute(
            select(SourceConnection).where(
                SourceConnection.user_id == current_user.id,
                SourceConnection.provider == "google_sheets",
                SourceConnection.source_url == source_url,
                SourceConnection.worksheet_name == sheet_name,
            )
        ).scalar_one_or_none()
        if existing_connection:
            existing_connection.connection_label = source_label
            existing_connection.mapping_profile_id = mapping_profile.id
            existing_connection.status = "connected"
            existing_connection.config_json = _json_dumps(parsed)
            existing_connection.updated_at = datetime.utcnow()
        else:
            db.add(
                SourceConnection(
                    user_id=current_user.id,
                    source_id=source.id,
                    provider="google_sheets",
                    connection_label=source_label,
                    source_url=source_url,
                    worksheet_name=sheet_name,
                    mapping_profile_id=mapping_profile.id,
                    status="connected",
                    config_json=_json_dumps(parsed),
                )
            )
    if upload_session:
        upload_session.status = "imported"
    db.commit()
    _log_audit_event(
        db,
        current_user.id,
        "shipment_imported",
        shipment_status="active",
        details={
            "imported_count": imported_count,
            "duplicate_count": duplicate_count,
            "skipped_blank_count": skipped_blank_count,
            "skipped_invalid_count": skipped_invalid_count,
            "mapping": mapping_json,
            "source_batch_id": batch.id if batch else 0,
            "source_type": source_type,
            "source_label": source_label,
            "mapping_profile_id": mapping_profile.id if mapping_profile else 0,
            "source_tracking_warning": source_tracking_warning,
        },
    )
    db.commit()
    try:
        file_path.unlink()
    except OSError as exc:
        logger.warning("Unable to delete temporary import file %s: %s", file_path, exc)
    return {
        "imported_count": imported_count,
        "duplicate_count": duplicate_count,
        "skipped_blank_count": skipped_blank_count,
        "skipped_invalid_count": skipped_invalid_count,
        "source_batch_id": batch.id if batch else 0,
        "source_label": source_label,
        "source_reference": source_reference,
        "mapping_profile_id": mapping_profile.id if mapping_profile else 0,
        "source_tracking_warning": source_tracking_warning,
        "total_rows": imported_count + duplicate_count + skipped_blank_count + skipped_invalid_count,
    }


@router.post("/import-validate")
async def import_validate(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    token = _clean_text(payload.get("temp_file_token"))
    mapping_json = payload.get("mapping_json") or {}
    row_overrides = payload.get("row_overrides") or []
    if not token:
        raise HTTPException(status_code=400, detail="temp_file_token is required")
    file_path = TEMP_IMPORT_DIR / token
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Temporary import file not found")

    customer_col = _clean_text(mapping_json.get("customer_name"))
    container_col = _clean_text(mapping_json.get("container_number"))
    bl_col = _clean_text(mapping_json.get("bl_number"))
    if not container_col:
        raise HTTPException(status_code=400, detail="Container Number mapping is required")

    upload_session = db.execute(
        select(UploadSession).where(
            UploadSession.user_id == current_user.id,
            UploadSession.stored_path == str(file_path),
        )
    ).scalar_one_or_none()
    selected_sheet = upload_session.detected_sheet if upload_session else None
    _sheet_name, header_row, headers, _preview_rows, rows, _available_sheets = _read_sheet(file_path, selected_sheet)
    header_index = header_row - 1
    normalized_rows = _normalize_import_row_objects(
        headers,
        rows,
        header_index,
        [customer_col, bl_col],
    )
    normalized_rows = _apply_import_row_overrides(
        normalized_rows,
        row_overrides,
        customer_col,
        container_col,
        bl_col,
    )
    review = _summarize_import_rows(normalized_rows, customer_col, container_col, bl_col)
    duplicates = _detect_import_duplicates(db, current_user, normalized_rows, customer_col, container_col, bl_col)
    return {
        "valid_count": review["valid_count"],
        "invalid_count": review["invalid_count"],
        "duplicate_count": duplicates["duplicate_count"],
        "skipped_blank_count": review["skipped_blank_count"],
        "invalid_rows": review["invalid_rows"],
        "duplicate_rows": duplicates["duplicate_rows"],
        "ready_to_import": review["invalid_count"] == 0,
    }


@router.post("/refresh-one")
def refresh_one(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    container_number = _clean_container(payload.get("container_number"))
    if not container_number:
        raise HTTPException(status_code=400, detail="container_number is required")
    shipments = list(
        db.execute(_user_shipment_select(current_user).where(Shipment.container_number == container_number)).scalars()
    )
    if not shipments:
        raise HTTPException(status_code=404, detail="Shipment not found")
    primary = shipments[0]
    _refresh_one_shipment(primary, use_cache=False)
    for shipment in shipments[1:]:
        shipment.latest_location = primary.latest_location
        shipment.latest_time = primary.latest_time
        shipment.train_no = primary.train_no
        shipment.departure = primary.departure
        shipment.rail_status = primary.rail_status
        shipment.movement_category = primary.movement_category
        shipment.delay_days = primary.delay_days
        shipment.wagon_no = primary.wagon_no
        shipment.train_origin = primary.train_origin
        shipment.train_destination = primary.train_destination
        shipment.shipping_line = primary.shipping_line
        shipment.tracking_source = primary.tracking_source
        shipment.last_refresh_at = primary.last_refresh_at
        shipment.last_refresh_status = primary.last_refresh_status
        shipment.last_refresh_error = primary.last_refresh_error
    db.commit()
    db.refresh(primary)
    return _shipment_to_dict(primary)


@router.post("/refresh-group")
def refresh_group(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    shipments = _resolve_group_shipments(
        db,
        current_user,
        bl_number=payload.get("bl_number"),
        container_numbers=payload.get("container_numbers") or [],
        include_archived=False,
    )
    if not shipments:
        raise HTTPException(status_code=404, detail="Shipment group not found")
    refreshed_group_count = len(_group_dashboard_rows(shipments))
    unique_containers = sorted({_clean_container(shipment.container_number) for shipment in shipments if _clean_container(shipment.container_number)})
    payloads: dict[str, dict[str, Any]] = {}
    with ThreadPoolExecutor(max_workers=max(1, min(TRACKING_POOL_WORKERS, len(unique_containers)))) as executor:
        futures = {executor.submit(_build_tracking_payload, container_number, False): container_number for container_number in unique_containers}
        for future in as_completed(futures):
            container_number = futures[future]
            try:
                payloads[container_number] = future.result()
            except Exception:
                logger.exception(
                    "Tracking payload build failed during group refresh for container %s",
                    container_number,
                )
                payloads[container_number] = {
                    "data": {},
                    "status": "error",
                    "error": "Tracking fetch failed",
                    "cached": False,
                    "has_data": False,
                }
    for shipment in shipments:
        container_number = _clean_container(shipment.container_number)
        if container_number in payloads:
            _apply_tracking_payload(shipment, payloads[container_number])
    _log_audit_event(
        db,
        current_user.id,
        "shipment_group_refreshed",
        bl_number=_first_non_empty([shipment.bl_number for shipment in shipments]),
        container_number=_first_non_empty([shipment.container_number for shipment in shipments]),
        shipment_status="active",
        details={
            "refreshed_count": refreshed_group_count,
            "refreshed_container_count": len(unique_containers),
            "container_numbers": unique_containers,
        },
    )
    db.commit()
    return {
        "refreshed_count": refreshed_group_count,
        "refreshed_container_count": len(unique_containers),
    }


@router.post("/refresh-all")
def refresh_all(db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    result = _refresh_active_shipments_for_user(
        db,
        current_user.id,
        audit_action="shipment_all_refreshed",
    )
    return {
        "refreshed_count": result["refreshed_count"],
        "refreshed_container_count": result["refreshed_container_count"],
        "message": f"Refreshed {result['refreshed_count']} active shipment{'' if result['refreshed_count'] == 1 else 's'}",
    }


@router.post("/refresh-all/start")
def start_refresh_all_tracking(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_portal_user),
):
    _ensure_user_scope_ready(db, current_user)
    active_shipments = list(
        db.execute(
            select(Shipment).where(
                Shipment.user_id == current_user.id,
                Shipment.shipment_status == "active",
            )
        ).scalars()
    )
    active_group_count = len(_group_dashboard_rows(active_shipments))
    unique_container_count = len(
        {
            _clean_container(shipment.container_number)
            for shipment in active_shipments
            if _clean_container(shipment.container_number)
        }
    )
    task_id = uuid4().hex
    _set_refresh_job(
        task_id,
        state="queued",
        total=unique_container_count,
        completed=0,
        progress=0,
        message=(
            f"Preparing live tracking refresh for {active_group_count} shipment"
            f"{'s' if active_group_count != 1 else ''}."
        ),
        user_id=current_user.id,
    )
    worker = Thread(
        target=_run_refresh_all_job,
        args=(task_id, normalized_database_url, current_user.id),
        daemon=True,
    )
    worker.start()
    return {"task_id": task_id}


@router.get("/refresh-all/status/{task_id}")
def get_refresh_all_tracking_status(
    task_id: str,
    current_user: User = Depends(get_portal_user),
):
    job = _get_refresh_job(task_id)
    if not job or job.get("user_id") != current_user.id:
        raise HTTPException(status_code=404, detail="Refresh task not found")
    return job


@router.post("/audit/group")
def get_group_audit_trail(payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    normalized_bl = _normalize_bl_number(payload.get("bl_number"))
    container_numbers = {
        _clean_container(value)
        for value in (payload.get("container_numbers") or [])
        if _clean_container(value)
    }
    query = select(AuditLog).where(AuditLog.user_id == current_user.id)
    entries = list(db.execute(query.order_by(AuditLog.created_at.desc(), AuditLog.id.desc())).scalars())
    filtered = []
    for entry in entries:
        entry_bl = _normalize_bl_number(entry.bl_number)
        entry_container = _clean_container(entry.container_number)
        details = _serialize_audit_log(entry).get("details", {})
        detail_containers = {
            _clean_container(value)
            for value in (details.get("container_numbers") or [])
            if _clean_container(value)
        }
        if normalized_bl and entry_bl == normalized_bl:
            filtered.append(entry)
            continue
        if container_numbers and (
            entry_container in container_numbers or bool(container_numbers.intersection(detail_containers))
        ):
            filtered.append(entry)
    return {"items": [_serialize_audit_log(entry) for entry in filtered[:100]]}


@router.post("/bl-documents/upload")
async def upload_bl_document(
    bl_number: str = Form(...),
    document_type: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_portal_user),
):
    _ensure_user_scope_ready(db, current_user)
    normalized_bl = _normalize_bl_number(bl_number)
    allowed = db.execute(
        _user_shipment_select(current_user).where(Shipment.bl_number == normalized_bl)
    ).scalars().first()
    if not allowed:
        raise HTTPException(status_code=404, detail="Shipment BL not found")
    metadata = _save_document(bl_number, document_type, file)
    _log_audit_event(
        db,
        current_user.id,
        "bl_document_uploaded",
        bl_number=normalized_bl,
        shipment_status=allowed.shipment_status,
        details={"document_type": document_type, "original_name": metadata.get("original_name", "")},
    )
    db.commit()
    return {"saved": True, "document": metadata}


@router.get("/bl-documents/file")
def get_bl_document_file(
    bl_number: str,
    document_type: str,
    db: Session = Depends(get_db),
    authorization: str | None = Header(default=None),
    access_token: str | None = None,
):
    _ensure_storage_ready(db)
    if access_token:
        current_user = _user_from_access_token(db, access_token)
    else:
        auth_text = _clean_text(authorization)
        if auth_text.lower().startswith("bearer "):
            current_user = _user_from_access_token(db, auth_text.split(" ", 1)[1])
        elif settings.allow_demo_portal_fallback:
            current_user = db.get(User, _resolve_default_user_id(db))
        else:
            raise HTTPException(status_code=401, detail="Missing authentication token")
    if not current_user:
        raise HTTPException(status_code=401, detail="User not found")
    _ensure_user_scope_ready(db, current_user)
    normalized_bl = _normalize_bl_number(bl_number)
    normalized_type = _clean_text(document_type).lower()
    if not normalized_bl or normalized_type not in VALID_DOCUMENT_TYPES:
        raise HTTPException(status_code=400, detail="Invalid document request")
    allowed = db.execute(
        _user_shipment_select(current_user).where(Shipment.bl_number == normalized_bl)
    ).scalars().first()
    if not allowed:
        raise HTTPException(status_code=404, detail="Document not found")
    metadata = (_load_documents_index().get(normalized_bl) or {}).get(normalized_type)
    if not isinstance(metadata, dict):
        raise HTTPException(status_code=404, detail="Document not found")
    storage_backend = _clean_text(metadata.get("storage_backend") or "local").lower()
    if storage_backend == "r2":
        object_key = _clean_text(metadata.get("object_key"))
        client = _get_r2_client()
        bucket_name = _r2_bucket_name()
        if client is None or not object_key:
            raise HTTPException(status_code=404, detail="Stored document file not found")
        try:
            response = client.get_object(Bucket=bucket_name, Key=object_key)
            content = response["Body"].read()
        except (BotoCoreError, ClientError, KeyError, OSError) as exc:
            raise HTTPException(status_code=404, detail="Stored document file not found") from exc
        filename = metadata.get("original_name") or Path(object_key).name
        return StreamingResponse(
            BytesIO(content),
            media_type=metadata.get("content_type") or "application/octet-stream",
            headers={"Content-Disposition": f'inline; filename="{filename}"'},
        )

    target_path = Path(metadata.get("path") or "")
    if not target_path.exists() or not target_path.is_file():
        raise HTTPException(status_code=404, detail="Stored document file not found")
    return FileResponse(
        path=target_path,
        media_type=metadata.get("content_type") or "application/octet-stream",
        filename=metadata.get("original_name") or target_path.name,
    )


@router.get("/bl-documents/open")
def open_bl_document_file(
    bl_number: str,
    document_type: str,
    db: Session = Depends(get_db),
    authorization: str | None = Header(default=None),
    access_token: str | None = None,
):
    return get_bl_document_file(bl_number, document_type, db, authorization, access_token)


@router.get("/cache/clear")
def clear_cache():
    if CACHE_FILE.exists():
        CACHE_FILE.unlink()
    return {"message": "Cache cleared successfully"}


@router.get("/cache/stats")
def cache_stats():
    if CACHE_FILE.exists():
        try:
            cache = json.loads(CACHE_FILE.read_text(encoding="utf-8"))
            return {"cache_size": len(cache), "cached_containers": list(cache.keys()), "cache_duration_minutes": CACHE_DURATION_MINUTES, "cache_file": str(CACHE_FILE)}
        except (OSError, json.JSONDecodeError, TypeError) as exc:
            logger.warning("Unable to read cache stats from %s: %s", CACHE_FILE, exc)
    return {"cache_size": 0, "cached_containers": [], "cache_duration_minutes": CACHE_DURATION_MINUTES}


@router.get("/test-container/{container_number}")
def test_container(container_number: str):
    container_number = _clean_container(container_number)
    ldb_result = _fetch_ldb(container_number)
    concor_result = _fetch_concor(container_number)
    return {"container": container_number, "ldb_api": {"success": ldb_result is not None, "data": ldb_result}, "concor_api": {"success": concor_result is not None, "data": concor_result}, "train_number_found": concor_result.get("train_no") if concor_result else None, "both_success": ldb_result is not None and concor_result is not None}


@router.get("/debug/concor-raw/{container_number}")
def debug_concor_raw(container_number: str):
    container_number = _clean_container(container_number)
    try:
        response = requests.post(CONCOR_API_URL, json={"containerNo": [container_number]}, timeout=20, headers={"Content-Type": "application/json", "Accept": "application/json, text/plain, */*", "User-Agent": "Mozilla/5.0", "Origin": "https://www.concorindia.co.in", "Referer": "https://www.concorindia.co.in/track-n-trace?lang=en"})
        return {"container": container_number, "status_code": response.status_code, "response_text": response.text[:2000] if response.text else "Empty", "response_json": response.json() if response.status_code == 200 else None}
    except (requests.RequestException, ValueError) as error:
        return {"container": container_number, "error": str(error)}


@router.get("/system/status")
def system_status(db: Session = Depends(get_db)):
    _ensure_storage_ready(db)
    total_shipments = db.execute(select(func.count()).select_from(Shipment)).scalar_one()
    active_shipments = db.execute(select(func.count()).select_from(Shipment).where(Shipment.shipment_status == "active")).scalar_one()
    customer_directory_count = db.execute(select(func.count()).select_from(CustomerDirectory)).scalar_one()
    return {"status": "running", "total_shipments": total_shipments, "active_shipments": active_shipments, "customer_directory_count": customer_directory_count, "apis_configured": {"ldb": LDB_API_URL, "concor": CONCOR_API_URL}, "cache_enabled": True, "cache_duration_minutes": CACHE_DURATION_MINUTES}


@router.get("/stats")
def get_stats(db: Session = Depends(get_db), current_user: User = Depends(get_portal_user)):
    _ensure_user_scope_ready(db, current_user)
    shipment_models = _get_shipments(db, current_user)
    shipments = [_shipment_to_dict(shipment) for shipment in shipment_models]
    grouped_rows = _group_dashboard_rows(shipment_models)
    movement_counts: dict[str, int] = {}
    refresh_status_counts: dict[str, int] = {}
    for shipment in grouped_rows:
        movement = shipment.get("movement_category") or "Unknown"
        movement_counts[movement] = movement_counts.get(movement, 0) + 1
    for shipment in shipments:
        refresh_status = shipment.get("last_refresh_status") or "not_refreshed"
        refresh_status_counts[refresh_status] = refresh_status_counts.get(refresh_status, 0) + 1
    active_shipments = [shipment for shipment in shipments if shipment.get("shipment_status") == "active"]
    return {
        "total_shipments": len(shipments),
        "active_shipments": len(active_shipments),
        "completed_shipments": len([s for s in shipments if s.get("shipment_status") == "completed"]),
        "archived_shipments": len([s for s in shipments if s.get("shipment_status") == "archived"]),
        "grouped_dashboard_rows": len(grouped_rows),
        "shipments_with_train_number": len([s for s in shipments if s.get("train_no")]),
        "shipments_with_location": len([s for s in shipments if s.get("latest_location")]),
        "movement_breakdown": movement_counts,
        "refresh_status_breakdown": refresh_status_counts,
        "last_refresh_times": [{"container_number": shipment.get("container_number", ""), "last_refresh_at": shipment.get("last_refresh_at", ""), "last_refresh_status": shipment.get("last_refresh_status", ""), "last_refresh_error": shipment.get("last_refresh_error", "")} for shipment in active_shipments],
    }
