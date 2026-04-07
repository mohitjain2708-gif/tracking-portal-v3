from __future__ import annotations
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

FIELD_ALIASES = {
    "container_number": ["container no", "container number", "container", "cntr no", "cntr", "container id"],
    "bl_number": ["bl", "bl no", "bill of lading", "bol"],
    "location": ["location", "current location", "yard location", "place", "station"],
    "status": ["status", "shipment status", "movement status"],
    "date": ["date", "updated date", "event date", "last updated"],
}

def normalize(value: Any) -> str:
    return str(value or "").strip().lower()

def score_sheet(ws) -> int:
    score = 0
    max_rows = min(ws.max_row, 15)
    max_cols = min(ws.max_column, 20)
    for r in range(1, max_rows + 1):
        for c in range(1, max_cols + 1):
            txt = normalize(ws.cell(r, c).value)
            if any(alias in txt for aliases in FIELD_ALIASES.values() for alias in aliases):
                score += 1
    return score

def detect_best_sheet(workbook) -> str:
    best_name = workbook.sheetnames[0]
    best_score = -1
    for name in workbook.sheetnames:
        score = score_sheet(workbook[name])
        if score > best_score:
            best_name = name
            best_score = score
    return best_name

def detect_header_row(ws) -> int:
    best_row = 1
    best_score = -1
    max_scan = min(ws.max_row, 25)
    for r in range(1, max_scan + 1):
        row_values = [normalize(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 50) + 1)]
        score = sum(
            1
            for cell in row_values
            if any(alias in cell for aliases in FIELD_ALIASES.values() for alias in aliases)
        )
        if score > best_score:
            best_row = r
            best_score = score
    if best_score < 1:
        raise ValueError("Could not detect a header row. Please upload a workbook with visible column titles.")
    return best_row

def extract_columns(ws, header_row: int) -> list[str]:
    columns = []
    for c in range(1, ws.max_column + 1):
        value = str(ws.cell(header_row, c).value or "").strip()
        if value:
            columns.append(value)
    return columns

def preview_rows(ws, header_row: int, limit: int = 10) -> list[dict]:
    headers = extract_columns(ws, header_row)
    rows = []
    for r in range(header_row + 1, min(ws.max_row, header_row + limit) + 1):
        item = {}
        empty = True
        for idx, header in enumerate(headers, start=1):
            val = ws.cell(r, idx).value
            if val not in (None, ""):
                empty = False
            item[header] = "" if val is None else str(val)
        if not empty:
            rows.append(item)
    return rows

def save_upload(file_bytes: bytes, destination: Path):
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(file_bytes)

def inspect_workbook(path: str) -> dict:
    workbook = load_workbook(path, data_only=True)
    sheet_name = detect_best_sheet(workbook)
    ws = workbook[sheet_name]
    header_row = detect_header_row(ws)
    columns = extract_columns(ws, header_row)
    previews = preview_rows(ws, header_row)
    return {
        "sheet_name": sheet_name,
        "header_row": header_row,
        "available_columns": columns,
        "preview_rows": previews,
    }

def process_with_mapping(path: str, sheet_name: str, header_row: int, mapping_json: dict) -> dict:
    workbook = load_workbook(path, data_only=True)
    ws = workbook[sheet_name]
    available_headers = {}
    for c in range(1, ws.max_column + 1):
        title = str(ws.cell(header_row, c).value or "").strip()
        if title:
            available_headers[title] = c

    output_rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        row_output = {}
        has_any = False
        for target_field, source_header in mapping_json.items():
            col_idx = available_headers.get(source_header)
            value = ws.cell(r, col_idx).value if col_idx else None
            if value not in (None, ""):
                has_any = True
            row_output[target_field] = "" if value is None else str(value)
        if has_any:
            output_rows.append(row_output)

    return {
        "headers": list(mapping_json.keys()),
        "rows": output_rows,
        "total_rows": len(output_rows),
    }
